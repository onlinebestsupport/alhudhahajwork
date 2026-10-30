# =================================================================================
# SECTION 14 (FLET 1.0.0 VERSION) — INVOICES TAB
# =================================================================================
# v4.5 — Ultimate fix for Flet 1.0.5 Dropdown sync & Layout.
#   • Restored fixed-width Container for AlertDialogs to fix squished desktop layout.
#   • Added an aggressive async polling loop (1500ms) in on_traveler_change to 
#     catch Flet's delayed dropdown value updates.
#   • Added fuzzy text matching to resolve IDs if Flet sends text.
# =================================================================================

import flet as ft
import json
import os
import time
import base64
import io
import asyncio
from pathlib import Path
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import pandas as pd

from core.helpers import (
    get_app_base_path,
    number_to_words_indian,
    format_currency_indian,
    send_file_to_user,
)
from core.settings_manager import SettingsManager
from core.company_settings_dialog import CompanySettingsDialog


TYPING_GUARD_SECONDS = 3.0


def _money(value):
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError, TypeError):
        return Decimal("0.00")


def _rate(value):
    try:
        return Decimal(str(value)).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError, TypeError):
        return Decimal("0.0000")


def _round_rule(value):
    value = Decimal(str(value))
    floor_value = value.to_integral_value(rounding="ROUND_FLOOR")
    return (value.to_integral_value(rounding="ROUND_CEILING")
            if value - floor_value >= Decimal("0.50") else floor_value)


# =================================================================================
# 14.1 — CLASS: InvoicesTab
# =================================================================================
class InvoicesTab:

    def __init__(self, page: ft.Page, db, current_user):
        self.page = page
        self.db = db
        self.current_user = current_user
        self.settings_manager = SettingsManager(db)

        self.invoices = []
        self.travelers = {}
        self.batches = {}
        self.traveler_details_cache = {}
        self.traveler_paid = {}
        self.invoiced_traveler_ids = set()

        self.stat_labels = {}
        self.tax_info_label = None
        self.status_filter = None
        self.search_input = None
        self.table = None
        self.root = None
        self.selection_label = None

        self._selected_invoice_id = None
        self._typing_until = 0.0

        self.setup_ui()
        self.refresh()

    def build(self):
        return self.root

    # =============================================================================
    # 14.1.0 — Helpers
    # =============================================================================
    def _get_effective_status(self, inv):
        status = str(inv.get('status', 'pending')).lower()
        due_date = self._fmt_date_ddmmyyyy(inv.get('due_date', ''))
        if status == 'pending' and due_date:
            try:
                due_dt = datetime.strptime(due_date, "%d/%m/%Y").date()
                if datetime.now().date() > due_dt:
                    return 'overdue'
            except Exception:
                pass
        return status

    def _safe_update(self):
        if time.time() < self._typing_until:
            return
        try:
            if self.root is not None:
                self.root.update()
        except Exception:
            pass

    def _update_table_only(self):
        for c in (self.table, self.selection_label):
            try:
                if c is not None:
                    c.update()
            except Exception:
                pass

    # =============================================================================
    # 14.1.1 — setup_ui
    # =============================================================================
    def setup_ui(self):
        tax = self.settings_manager.get_tax_settings()
        try:
            gst_rate = _rate(tax.get('gst_percentage', 18))
            tcs_rate = _rate(tax.get('tcs_percentage', 0.1))
        except Exception:
            gst_rate = Decimal('18')
            tcs_rate = Decimal('0.1')

        self.tax_info_label = ft.Text(
            f"GST {gst_rate}%  |  TCS {tcs_rate}%",
            size=11, weight=ft.FontWeight.BOLD,
            color="#FFD700",
            no_wrap=False, max_lines=2)

        def refresh_tax(e):
            self.update_tax_display()
            self._snack("✅ Tax rates refreshed")

        def open_tax_settings(e):
            dlg = CompanySettingsDialog(
                self.page, self.db, self.current_user,
                on_save_callback=self.refresh)
            dlg.show()

        tax_bar = ft.Container(
            content=ft.ResponsiveRow(
                controls=[
                    ft.Container(
                        content=ft.Row([
                            ft.Text("💰", size=18),
                            self.tax_info_label,
                        ], spacing=6),
                        col={"xs": 12, "sm": 12, "md": 6}),
                    ft.Container(
                        content=ft.Row([
                            ft.Button(
                                content=ft.Text("🔄 Refresh",
                                                size=10,
                                                weight=ft.FontWeight.BOLD,
                                                color=ft.Colors.WHITE),
                                on_click=refresh_tax, height=34,
                                bgcolor="#f39c12"),
                            ft.Button(
                                content=ft.Text("⚙️ Settings",
                                                size=10,
                                                weight=ft.FontWeight.BOLD,
                                                color=ft.Colors.WHITE),
                                on_click=open_tax_settings, height=34,
                                bgcolor="#2ecc71"),
                        ], spacing=6, alignment=ft.MainAxisAlignment.END),
                        col={"xs": 12, "sm": 12, "md": 6}),
                ],
                spacing=8, run_spacing=8,
            ),
            padding=ft.Padding.symmetric(horizontal=12, vertical=10),
            gradient=ft.LinearGradient(
                begin=ft.Alignment.CENTER_LEFT,
                end=ft.Alignment.CENTER_RIGHT,
                colors=["#2c3e50", "#1a252f", "#2c3e50"]),
            border_radius=12,
        )

        stat_configs = [
            ("total_invoices",   "📊 Invoices",   "#3498db"),
            ("pending",          "⏳ Pending",    "#f39c12"),
            ("paid",             "✅ Paid",       "#27ae60"),
            ("discount",         "🎁 Discount",   "#c2185b"),
            ("gst",              "🏛️ GST",        "#e67e22"),
            ("tcs",              "💰 TCS",        "#9b59b6"),
            ("package_pending",  "📦 Pkg Pend",   "#e74c3c"),
        ]

        stat_cards = []
        for key, label, color in stat_configs:
            value_label = ft.Text("0", size=14,
                                  weight=ft.FontWeight.BOLD,
                                  color=ft.Colors.WHITE)
            self.stat_labels[key] = value_label
            card = ft.Container(
                content=ft.Column(
                    controls=[
                        ft.Text(label, size=9, color=ft.Colors.WHITE,
                                weight=ft.FontWeight.BOLD,
                                no_wrap=False, max_lines=2),
                        value_label,
                    ],
                    spacing=2,
                    horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                ),
                padding=8,
                bgcolor=color,
                border_radius=10,
                height=64,
            )
            stat_cards.append(card)

        stats_row = ft.ResponsiveRow(
            controls=[
                ft.Container(content=c, col={"xs": 6, "sm": 4, "md": 3, "lg": 3})
                for c in stat_cards
            ],
            spacing=6, run_spacing=6,
        )

        def _tb(label, color, handler):
            return ft.Button(
                content=ft.Text(label, size=11,
                                weight=ft.FontWeight.BOLD,
                                color=ft.Colors.WHITE,
                                no_wrap=True,
                                overflow=ft.TextOverflow.ELLIPSIS),
                on_click=handler, height=40,
                bgcolor=color,
                style=ft.ButtonStyle(
                    shape=ft.RoundedRectangleBorder(radius=8)))

        self.status_filter = ft.Dropdown(
            label="Status",
            options=[
                ft.dropdown.Option(key="All", text="All"),
                ft.dropdown.Option(key="pending", text="Pending"),
                ft.dropdown.Option(key="paid", text="Paid"),
                ft.dropdown.Option(key="overdue", text="Overdue"),
                ft.dropdown.Option(key="cancelled", text="Cancelled"),
            ],
            value="All", height=48, text_size=11)
        self.status_filter.on_change = self.apply_filters

        self.search_input = ft.TextField(
            hint_text="🔍 Search invoice or traveler...",
            height=48, text_size=12,
            content_padding=ft.Padding.symmetric(horizontal=10, vertical=8))

        def _search_on_change(e, _f=self.search_input):
            self._typing_until = time.time() + TYPING_GUARD_SECONDS
            try:
                ctrl = getattr(e, "control", None)
                if ctrl is not None and hasattr(ctrl, "value"):
                    if _f.value != ctrl.value:
                        _f.value = ctrl.value
            except Exception:
                pass
            self.apply_filters(e)

        self.search_input.on_change = _search_on_change

        toolbar = ft.ResponsiveRow(
            controls=[
                ft.Container(
                    content=_tb("➕ Generate Invoice", "#27ae60",
                                self.generate_invoice),
                    col={"xs": 6, "sm": 4, "md": 3}),
                ft.Container(
                    content=_tb("🧾 Manual Create", "#f39c12",
                                self.open_manual_invoice),
                    col={"xs": 6, "sm": 4, "md": 3}),
                ft.Container(
                    content=_tb("🧹 Cleanup Duplicates", "#7f8c8d",
                                self.open_payment_cleanup),
                    col={"xs": 6, "sm": 4, "md": 3}),
                ft.Container(
                    content=_tb("📄 PDF (selected)", "#e74c3c",
                                self.export_pdf_selected),
                    col={"xs": 6, "sm": 4, "md": 2}),
                ft.Container(
                    content=_tb("🖨️ Print (selected)", "#9b59b6",
                                self.print_invoice_selected),
                    col={"xs": 6, "sm": 4, "md": 2}),
                ft.Container(content=self.status_filter,
                             col={"xs": 6, "sm": 4, "md": 2}),
                ft.Container(content=self.search_input,
                             col={"xs": 12, "sm": 12, "md": 4}),
            ],
            spacing=8, run_spacing=8,
        )

        self.table = ft.DataTable(
            columns=[
                ft.DataColumn(ft.Text("Sel", size=11)),
                ft.DataColumn(ft.Text("⋮", size=11)),
                ft.DataColumn(ft.Text("Invoice", size=11)),
                ft.DataColumn(ft.Text("Date", size=11)),
                ft.DataColumn(ft.Text("Traveler", size=11)),
                ft.DataColumn(ft.Text("Base", size=11)),
                ft.DataColumn(ft.Text("Disc%", size=11)),
                ft.DataColumn(ft.Text("Disc", size=11)),
                ft.DataColumn(ft.Text("Taxable", size=11)),
                ft.DataColumn(ft.Text("GST%", size=11)),
                ft.DataColumn(ft.Text("GST", size=11)),
                ft.DataColumn(ft.Text("TCS%", size=11)),
                ft.DataColumn(ft.Text("TCS", size=11)),
                ft.DataColumn(ft.Text("Total", size=11)),
                ft.DataColumn(ft.Text("Rounded", size=11)),
                ft.DataColumn(ft.Text("Due", size=11)),
                ft.DataColumn(ft.Text("Status", size=11)),
                ft.DataColumn(ft.Text("Pkg Pend", size=11)),
            ],
            rows=[],
            column_spacing=10,
            heading_row_color=ft.Colors.BLUE_GREY_800,
            heading_row_height=40,
            data_row_min_height=44,
            data_row_max_height=56,
            border=ft.Border.all(1, ft.Colors.GREY_300),
            border_radius=10,
            vertical_lines=ft.BorderSide(1, ft.Colors.GREY_200),
            horizontal_lines=ft.BorderSide(1, ft.Colors.GREY_200),
        )

        self.selection_label = ft.Text("", size=10,
                                       color=ft.Colors.BLUE_700,
                                       weight=ft.FontWeight.BOLD,
                                       italic=True)

        self.root = ft.Container(
            content=ft.Column(
                controls=[
                    tax_bar,
                    stats_row,
                    ft.Container(content=toolbar, padding=8,
                                 bgcolor=ft.Colors.WHITE,
                                 border_radius=10),
                    ft.Container(
                        content=ft.Column([
                            ft.Text("💡 Tap a row to SELECT it, then use "
                                    "toolbar 📄 PDF / 🖨️ Print. "
                                    "Use ⋮ menu for per-row actions.",
                                    size=10,
                                    color=ft.Colors.GREY_600,
                                    italic=True),
                            self.selection_label,
                            ft.Row([self.table],
                                   scroll=ft.ScrollMode.ADAPTIVE),
                        ], spacing=6),
                        bgcolor=ft.Colors.WHITE,
                        border_radius=10,
                        padding=10),
                ],
                spacing=10,
                scroll=ft.ScrollMode.AUTO,
            ),
            padding=10, bgcolor="#f0f2f5", expand=True,
        )

    def _darken(self, color):
        return {
            "#3498db": "#2471a3", "#27ae60": "#1e8449",
            "#f39c12": "#d68910", "#9b59b6": "#7d3c98",
            "#e74c3c": "#c0392b", "#c2185b": "#880e4f",
        }.get(color, color)

    # =============================================================================
    # 14.1.2 — Helpers
    # =============================================================================
    def safe_float(self, value, default=Decimal("0")):
        if value is None:
            return default
        text = str(value).strip()
        if not text or text.lower() in {"nan", "none", "nat", "null"}:
            return default
        try:
            return Decimal(text)
        except (InvalidOperation, ValueError, TypeError):
            return default

    def safe_str(self, value, default=''):
        if value is None:
            return default
        if isinstance(value, float):
            if value != value or str(value) == 'nan':
                return default
            if value.is_integer():
                return str(int(value))
            return str(value)
        return str(value)

    def _fmt_date_ddmmyyyy(self, date_val):
        if date_val is None:
            return ''
        s = str(date_val).strip()
        if not s or s.lower() in ('nan', 'none', 'nat', ''):
            return ''
        if len(s) >= 10 and s[4] == '-' and s[7] == '-':
            return f"{s[8:10]}/{s[5:7]}/{s[0:4]}"
        for fmt in ('%Y-%m-%dT%H:%M:%S', '%Y-%m-%d %H:%M:%S',
                    '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y'):
            try:
                return datetime.strptime(s[:19], fmt).strftime('%d/%m/%Y')
            except Exception:
                continue
        return s[:10]

    def update_tax_display(self):
        tax = self.settings_manager.get_tax_settings()
        try:
            gst_rate = _rate(tax.get('gst_percentage', 18))
            tcs_rate = _rate(tax.get('tcs_percentage', 0.1))
        except Exception:
            gst_rate = Decimal('18')
            tcs_rate = Decimal('0.1')
        self.tax_info_label.value = (
            f"GST {gst_rate}%  |  TCS {tcs_rate}%")
        try:
            self.tax_info_label.update()
        except Exception:
            pass

    # =============================================================================
    # 14.1.3 — refresh
    # =============================================================================
    def refresh(self, e=None):
        try:
            self.invoices = self.db.get_invoices()
            self.invoiced_traveler_ids = {
                str(i.get('traveler_id') or '').strip()
                for i in self.invoices
                if str(i.get('traveler_id') or '').strip()
            }
            self.travelers = {
                str(t.get('id') or '').strip(): (
                    f"{t.get('first_name', '')} "
                    f"{t.get('last_name', '')}"
                ).strip() or "Unnamed"
                for t in self.db.get_travelers()
                if str(t.get('id') or '').strip()
            }
            self.batches = {
                str(b.get('id') or '').strip(): b
                for b in self.db.get_batches()
                if str(b.get('id') or '').strip()
            }

            self.traveler_details_cache = {}
            for t in self.db.get_travelers():
                tid = t.get('id')
                if tid:
                    tid = str(t.get('id') or '').strip()
                    batch_id = str(t.get('batch_id') or '').strip()
                    bp = 0
                    bn = 'N/A'
                    if batch_id and batch_id in self.batches:
                        bp = _money(self.batches[batch_id].get('price', 0) or 0)
                        bn = self.batches[batch_id].get('batch_name', 'N/A')
                    self.traveler_details_cache[tid] = {
                        'name': (f"{t.get('first_name', '')} "
                                 f"{t.get('last_name', '')}").strip()
                                or 'Unknown',
                        'batch_id': batch_id,
                        'batch_price': bp,
                        'batch_name': bn,
                    }

            self.traveler_paid = {}
            for p in self.db.get_payments():
                tid = str(p.get('traveler_id') or '').strip()
                status = str(p.get('status', '')).strip().lower()
                if tid and status == 'completed':
                    self.traveler_paid[tid] = (
                        self.traveler_paid.get(tid, 0)
                        + _money(p.get('amount', 0) or 0))

            print(f"[INVOICES] loaded {len(self.invoices)} invoices, "
                  f"{len(self.travelers)} travelers")

            self.display_invoices()
            self.update_summary_stats()
            self.update_tax_display()
            self._safe_update()
        except Exception as ex:
            print(f"Invoices refresh error: {ex}")
            import traceback
            traceback.print_exc()

    # =============================================================================
    # 14.1.4 — display_invoices
    # =============================================================================
    def display_invoices(self, invoices=None):
        if invoices is None:
            invoices = self.invoices

        self.table.rows.clear()
        for inv in invoices:
            inv_id = inv.get('id')
            inv_no = self.safe_str(inv.get('invoice_no', ''))
            date_str = self._fmt_date_ddmmyyyy(inv.get('issue_date', ''))
            traveler_name = self.travelers.get(
                inv.get('traveler_id', ''), 'Unknown')

            base = self.safe_float(inv.get('amount', 0))
            disc_pct = self.safe_float(inv.get('discount_percentage', 0))
            disc_amt = self.safe_float(inv.get('discount_amount', 0))
            taxable = self.safe_float(
                inv.get('taxable_value', base - disc_amt))
            gst_pct = self.safe_float(inv.get('gst_percentage', 18))
            gst_amt = self.safe_float(inv.get('gst_amount', 0))
            tcs_pct = self.safe_float(inv.get('tcs_percentage', 0.1))
            tcs_amt = self.safe_float(inv.get('tcs_amount', 0))
            total = self.safe_float(inv.get('total_amount', 0))
            rounded = self.safe_float(
                inv.get('rounded_total', total))
            due_date = self._fmt_date_ddmmyyyy(inv.get('due_date', ''))
            
            eff_status = self._get_effective_status(inv)
            display_status = eff_status.upper()
            status_color = ("#27ae60" if eff_status == "paid"
                            else "#f39c12" if eff_status == "pending"
                            else "#e74c3c" if eff_status == "overdue"
                            else "#95a5a6")

            tid = str(inv.get('traveler_id') or '').strip()
            tdet = self.traveler_details_cache.get(
                tid, {'batch_price': 0})
            bp = tdet.get('batch_price', 0)
            paid = self.traveler_paid.get(tid, 0)
            
            pkg_pending = (
                max(Decimal('0.00'), bp - paid)
                if tid in self.invoiced_traveler_ids
                else Decimal('0.00'))
            
            if bp > 0:
                if pkg_pending <= 0:
                    pkg_txt = "✅ Paid"
                    pkg_color = "#27ae60"
                else:
                    pkg_txt = f"₹{pkg_pending:,.0f}"
                    pkg_color = ("#e67e22" if pkg_pending < bp * Decimal('0.5')
                                 else "#e74c3c")
            else:
                pkg_txt = "N/A"
                pkg_color = "#95a5a6"

            def _make_actions_menu(_inv_id=inv_id):
                def _fresh():
                    return next(
                        (x for x in self.invoices
                         if x.get('id') == _inv_id), None)

                def _wrap(handler):
                    def h(e, _h=handler, _f=_fresh):
                        fresh = _f()
                        if fresh is None:
                            self._snack("⚠️ Invoice no longer exists — "
                                        "refreshing…")
                            self.refresh()
                            return
                        _h(fresh)
                    return h

                return ft.PopupMenuButton(
                    icon=ft.Icons.MORE_VERT,
                    icon_color="#3498db",
                    icon_size=20,
                    tooltip="Invoice Actions",
                    items=[
                        ft.PopupMenuItem(
                            content=ft.Text("👁️  View Details"),
                            on_click=_wrap(self.view_invoice_details)),
                        ft.PopupMenuItem(
                            content=ft.Text("✏️  Modify Invoice"),
                            on_click=_wrap(self.modify_invoice)),
                        ft.PopupMenuItem(),
                        ft.PopupMenuItem(
                            content=ft.Text("📄  Export PDF"),
                            on_click=_wrap(self.export_invoice_to_pdf)),
                        ft.PopupMenuItem(
                            content=ft.Text("📊  Export Excel"),
                            on_click=_wrap(self.export_invoice_to_excel)),
                        ft.PopupMenuItem(
                            content=ft.Text("🖨️  Print Invoice"),
                            on_click=_wrap(self.print_invoice)),
                        ft.PopupMenuItem(),
                        ft.PopupMenuItem(
                            content=ft.Text("🗑️  Delete"),
                            on_click=_wrap(self.delete_invoice)),
                    ],
                )

            is_selected = (self._selected_invoice_id == inv_id)

            def _on_row_tap(e, _inv_id=inv_id):
                try:
                    self._select_invoice(_inv_id)
                except Exception as ex:
                    print(f"[INVOICES] row tap error: {ex}")

            self.table.rows.append(
                ft.DataRow(
                    on_select_change=_on_row_tap,
                    selected=is_selected,
                    cells=[
                        ft.DataCell(
                            ft.Container(
                                content=ft.Text(
                                    "✓" if is_selected else "",
                                    size=14,
                                    weight=ft.FontWeight.BOLD,
                                    color="#27ae60"),
                                width=24,
                                alignment=ft.Alignment.CENTER,
                                bgcolor=("#dcfce7" if is_selected
                                         else None),
                                border_radius=4,
                            )),
                        ft.DataCell(_make_actions_menu()),
                        ft.DataCell(ft.Text(inv_no[:16], size=10,
                                            weight=ft.FontWeight.BOLD)),
                        ft.DataCell(ft.Text(date_str, size=10)),
                        ft.DataCell(ft.Text(traveler_name[:18], size=10)),
                        ft.DataCell(ft.Text(f"₹{base:,.0f}", size=10)),
                        ft.DataCell(ft.Text(
                            f"{disc_pct:.1f}%", size=10,
                            color="#c2185b" if disc_pct > 0 else None)),
                        ft.DataCell(ft.Text(
                            f"₹{disc_amt:,.0f}", size=10,
                            color="#c2185b" if disc_amt > 0 else None)),
                        ft.DataCell(ft.Text(
                            f"₹{taxable:,.0f}", size=10,
                            weight=ft.FontWeight.BOLD)),
                        ft.DataCell(ft.Text(f"{gst_pct}%", size=10)),
                        ft.DataCell(ft.Text(f"₹{gst_amt:,.0f}", size=10)),
                        ft.DataCell(ft.Text(f"{tcs_pct}%", size=10)),
                        ft.DataCell(ft.Text(f"₹{tcs_amt:,.0f}", size=10)),
                        ft.DataCell(ft.Text(f"₹{total:,.0f}", size=10,
                                            weight=ft.FontWeight.BOLD)),
                        ft.DataCell(ft.Text(f"₹{rounded:,.0f}", size=10,
                                            color="#003366",
                                            weight=ft.FontWeight.BOLD)),
                        ft.DataCell(ft.Text(due_date, size=10)),
                        ft.DataCell(ft.Text(display_status, size=10,
                                            color=status_color,
                                            weight=ft.FontWeight.BOLD)),
                        ft.DataCell(ft.Text(pkg_txt, size=10,
                                            color=pkg_color)),
                    ]))

    # =============================================================================
    # 14.1.4b — _select_invoice / _get_selected_invoice
    # =============================================================================
    def _select_invoice(self, invoice_id):
        if self._selected_invoice_id == invoice_id:
            self._selected_invoice_id = None
        else:
            self._selected_invoice_id = invoice_id

        if self.selection_label:
            if self._selected_invoice_id:
                inv = next((x for x in self.invoices
                            if x.get('id') == self._selected_invoice_id),
                           None)
                if inv:
                    tname = self.travelers.get(
                        inv.get('traveler_id', ''), '?')
                    self.selection_label.value = (
                        f"✅ Selected: "
                        f"{self.safe_str(inv.get('invoice_no', ''))[:16]} | "
                        f"{tname} | "
                        f"₹{self.safe_float(inv.get('rounded_total', 0)):,.0f}")
                else:
                    self.selection_label.value = ""
            else:
                self.selection_label.value = ""

        self.display_invoices()
        self._update_table_only()

    def _get_selected_invoice(self):
        if not self._selected_invoice_id:
            return None
        return next((x for x in self.invoices
                     if x.get('id') == self._selected_invoice_id), None)

    # =============================================================================
    # 14.1.5 — Stats
    # =============================================================================
    def update_summary_stats(self):
        total_inv = len(self.invoices)
        
        pending = sum(self.safe_float(i.get('rounded_total', i.get('total_amount', 0)))
                      for i in self.invoices
                      if self._get_effective_status(i) == 'pending')
        paid = sum(self.safe_float(i.get('rounded_total', i.get('total_amount', 0)))
                   for i in self.invoices
                   if self._get_effective_status(i) == 'paid')
        discount = sum(self.safe_float(i.get('discount_amount', 0))
                       for i in self.invoices)
        gst = sum(self.safe_float(i.get('gst_amount', 0))
                  for i in self.invoices)
        tcs = sum(self.safe_float(i.get('tcs_amount', 0))
                  for i in self.invoices)

        pkg_pending = Decimal('0.00')
        for tid, paid_amt in self.traveler_paid.items():
            if tid not in self.invoiced_traveler_ids:
                continue
            det = self.traveler_details_cache.get(
                tid, {'batch_price': 0})
            bp = det.get('batch_price', 0)
            if bp > 0 and bp - paid_amt > 0:
                pkg_pending += bp - paid_amt

        def _short(v):
            try:
                n = Decimal(str(v))
            except Exception:
                return "₹0"
            if n >= 10000000:
                return f"₹{n/Decimal('10000000'):.2f}Cr"
            if n >= 100000:
                return f"₹{n/Decimal('100000'):.2f}L"
            if n >= 1000:
                return f"₹{n/Decimal('1000'):.1f}K"
            return f"₹{n:,.0f}"

        self.stat_labels['total_invoices'].value = str(total_inv)
        self.stat_labels['pending'].value = _short(pending)
        self.stat_labels['paid'].value = _short(paid)
        self.stat_labels['discount'].value = _short(discount)
        self.stat_labels['gst'].value = _short(gst)
        self.stat_labels['tcs'].value = _short(tcs)
        self.stat_labels['package_pending'].value = _short(pkg_pending)

    # =============================================================================
    # 14.1.6 — Filter
    # =============================================================================
    def _apply_filters_internal(self):
        status = (self.status_filter.value or "All") \
            if self.status_filter else "All"
        search = (self.search_input.value or "").lower().strip() \
            if self.search_input else ""
        filtered = self.invoices
        
        if status != "All":
            filtered = [i for i in filtered if self._get_effective_status(i) == status]
            
        if search:
            filtered = [
                i for i in filtered
                if (search in str(i.get('invoice_no', '')).lower()
                    or search in self.travelers.get(
                        i.get('traveler_id', ''), '').lower())]
        self.display_invoices(filtered)

    def apply_filters(self, e=None):
        self._apply_filters_internal()
        self._update_table_only()

    # =============================================================================
    # 14.1.7 — CRUD
    # =============================================================================
    def generate_invoice(self, e):
        dlg = InvoiceDialog(self.page, self.db, self.current_user,
                            on_save=self.refresh)
        dlg.show()

    def modify_invoice(self, invoice):
        dlg = InvoiceModifyDialog(self.page, self.db, self.current_user,
                                  invoice, on_save=self.refresh)
        dlg.show()

    def view_invoice_details(self, invoice):
        dlg = InvoiceDetailsDialog(self.page, self.db, invoice)
        dlg.show()

    def delete_invoice(self, invoice):
        def confirm(ev):
            try:
                self.db.delete_invoice(invoice['id'])
                self.db.log_activity(
                    self.current_user['id'], "delete_invoice",
                    f"Deleted invoice {invoice.get('invoice_no', '')}; "
                    "payments and receipts preserved")
                self.page.pop_dialog()
                self.refresh()
                self._snack("✅ Invoice deleted; payments and receipts kept")
            except Exception as ex:
                self._snack(f"❌ {ex}")

        dialog = ft.AlertDialog(
            title=ft.Text("Delete Invoice?", size=14),
            content=ft.Text(
                f"Delete invoice {invoice.get('invoice_no', 'N/A')}?\n\n"
                "Only the invoice will be deleted. Linked payments and "
                "receipts will be preserved for verification.",
                size=12),
            actions=[
                ft.TextButton(content=ft.Text("Cancel"),
                              on_click=lambda e: self.page.pop_dialog()),
                ft.Button(content=ft.Text("Delete"), on_click=confirm,
                          bgcolor=ft.Colors.RED_600,
                          color=ft.Colors.WHITE),
            ],
        )
        self.page.show_dialog(dialog)

    def open_manual_invoice(self, e):
        dlg = ManualInvoiceDialog(self.page, self.db, self.current_user,
                                  on_save=self.refresh)
        dlg.show()

    def open_payment_cleanup(self, e=None):
        """Show suspicious duplicate/automatic payment records for review."""
        try:
            payments = list(self.db.get_payments() or [])
        except Exception as ex:
            self._snack(f"❌ Could not load payments: {ex}")
            return

        candidates = []
        signatures = {}
        for index, payment in enumerate(payments):
            method = str(payment.get('payment_method', '') or '').strip().lower()
            notes = str(payment.get('notes', '') or '').strip().lower()
            is_auto = ('auto (invoice status)' in method
                       or 'auto-created when invoice' in notes)
            amount = _money(payment.get('amount', 0) or 0)
            signature = (
                str(payment.get('traveler_id', '') or '').strip(),
                str(amount),
                str(payment.get('payment_date', '') or '').strip(),
                method,
                str(payment.get('status', '') or '').strip().lower(),
                str(payment.get('invoice_id', '') or '').strip(),
            )
            signatures.setdefault(signature, []).append(index)
            if is_auto:
                candidates.append((index, payment, "automatic invoice payment"))

        for indexes in signatures.values():
            if len(indexes) > 1:
                for index in indexes[1:]:
                    item = (index, payments[index], "exact duplicate payment")
                    if not any(x[0] == index for x in candidates):
                        candidates.append(item)

        if not candidates:
            self._snack("✅ No suspicious duplicate or automatic payments found")
            return

        checks = []
        for index, payment, reason in candidates:
            pid = str(payment.get('id', '') or '').strip() or f"row {index + 2}"
            date = str(payment.get('payment_date', '') or '')[:10] or '—'
            traveler = str(payment.get('traveler_id', '') or '—')
            amount = _money(payment.get('amount', 0) or 0)
            checks.append(ft.Checkbox(
                label=(f"{reason}: {pid} | {date} | {traveler} | "
                       f"₹{amount:,.2f}"),
                value=False))

        def confirm_cleanup(ev):
            selected = [
                candidates[i][1].get('id')
                for i, checkbox in enumerate(checks)
                if checkbox.value and candidates[i][1].get('id')
            ]
            if not selected:
                self._snack("⚠️ Check at least one record to delete")
                return
            try:
                result = self.db.delete_payment_records(selected)
                self.page.pop_dialog()
                self.refresh()
                self._snack(
                    f"✅ Removed {result.get('payments_deleted', 0)} payment(s) "
                    f"and {result.get('receipts_deleted', 0)} receipt(s)")
            except Exception as ex:
                self._snack(f"❌ Cleanup failed: {ex}")

        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("🧹 Review Duplicate Payments", size=14),
            content=ft.Container(
                content=ft.Column([
                    ft.Text(
                        "Select only records you confirm are duplicates or old "
                        "automatic invoice payments. Checked records and their "
                        "receipts will be deleted.", size=11),
                    ft.Column(checks, spacing=4,
                              scroll=ft.ScrollMode.AUTO),
                ], spacing=10, scroll=ft.ScrollMode.AUTO),
                width=650, height=460, padding=10),
            actions=[
                ft.TextButton(content=ft.Text("Cancel"),
                              on_click=lambda ev: self.page.pop_dialog()),
                ft.Button(content=ft.Text("Delete Checked"),
                          on_click=confirm_cleanup,
                          bgcolor=ft.Colors.RED_600,
                          color=ft.Colors.WHITE),
            ],
        )
        self.page.show_dialog(dialog)

    # =============================================================================
    # 14.1.8 — Toolbar PDF / Print
    # =============================================================================
    def export_pdf_selected(self, e):
        inv = self._get_selected_invoice()
        if not inv:
            self._snack("⚠️ Tap a row first to select it, then tap 📄 PDF")
            return
        self.export_invoice_to_pdf(inv)

    def print_invoice_selected(self, e):
        inv = self._get_selected_invoice()
        if not inv:
            self._snack("⚠️ Tap a row first to select it, then tap 🖨️ Print")
            return
        self.print_invoice(inv)

    # =============================================================================
    # 14.1.8b — export_invoice_to_pdf
    # =============================================================================
    def export_invoice_to_pdf(self, invoice):
        try:
            data = self._get_invoice_data(invoice)
            from reportlab.lib.pagesizes import A4
            from reportlab.lib import colors as rl_colors
            from reportlab.platypus import (
                SimpleDocTemplate, Table, TableStyle, Paragraph,
                Spacer, Image)
            from reportlab.lib.styles import (
                getSampleStyleSheet, ParagraphStyle)
            from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
            from reportlab.lib.units import cm, inch

            base = get_app_base_path()
            out_dir = Path(base) / "invoices"
            out_dir.mkdir(exist_ok=True)
            fname = (f"invoice_{data['invoice_no']}_"
                     f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")
            filepath = out_dir / fname

            doc = SimpleDocTemplate(
                str(filepath), pagesize=A4,
                topMargin=0.5 * cm, bottomMargin=0.5 * cm,
                leftMargin=0.6 * cm, rightMargin=0.6 * cm)
            elements = []
            styles = getSampleStyleSheet()

            company_name_style = ParagraphStyle(
                'CN', parent=styles['Heading1'], fontSize=15,
                fontName='Helvetica-Bold', alignment=TA_RIGHT,
                spaceAfter=1, leading=17,
                textColor=rl_colors.HexColor('#1a252f'))
            company_info_style = ParagraphStyle(
                'CI', parent=styles['Normal'], fontSize=8,
                alignment=TA_RIGHT, spaceAfter=0, leading=10,
                textColor=rl_colors.HexColor('#555555'))
            title_style = ParagraphStyle(
                'T', parent=styles['Heading1'], fontSize=14,
                fontName='Helvetica-Bold', alignment=TA_CENTER,
                spaceBefore=4, spaceAfter=4, leading=16,
                textColor=rl_colors.HexColor('#2c3e50'))
            meta_l = ParagraphStyle('ML', parent=styles['Normal'],
                                    fontSize=10, fontName='Helvetica-Bold',
                                    alignment=TA_LEFT, leading=12)
            meta_r = ParagraphStyle('MR', parent=styles['Normal'],
                                    fontSize=10, fontName='Helvetica-Bold',
                                    alignment=TA_RIGHT, leading=12)
            box_lbl = ParagraphStyle('BL', parent=styles['Normal'],
                                     fontSize=8, fontName='Helvetica-Bold',
                                     alignment=TA_CENTER, leading=10,
                                     textColor=rl_colors.HexColor('#333'))
            body_l = ParagraphStyle('BdL', parent=styles['Normal'],
                                    fontSize=8, alignment=TA_LEFT, leading=10)
            body_r = ParagraphStyle('BdR', parent=styles['Normal'],
                                    fontSize=8, alignment=TA_RIGHT, leading=10)
            cell = ParagraphStyle('Ce', parent=styles['Normal'],
                                  fontSize=6.5, alignment=TA_LEFT, leading=8)
            cell_c = ParagraphStyle('Cc', parent=cell, alignment=TA_CENTER)
            cell_r = ParagraphStyle('Cr', parent=cell, alignment=TA_RIGHT)
            hdr = ParagraphStyle('H', parent=styles['Normal'], fontSize=6.5,
                                 fontName='Helvetica-Bold',
                                 alignment=TA_CENTER, leading=8,
                                 textColor=rl_colors.whitesmoke)
            t_lbl = ParagraphStyle('TL', parent=styles['Normal'],
                                   fontSize=8, alignment=TA_RIGHT, leading=10)
            t_val = ParagraphStyle('TV', parent=styles['Normal'],
                                   fontSize=8, alignment=TA_RIGHT, leading=10)
            t_disc_l = ParagraphStyle('TDL', parent=t_lbl,
                                      textColor=rl_colors.HexColor('#c2185b'),
                                      fontName='Helvetica-Bold')
            t_disc_v = ParagraphStyle('TDV', parent=t_val,
                                      textColor=rl_colors.HexColor('#c2185b'),
                                      fontName='Helvetica-Bold')
            g_lbl = ParagraphStyle('GL', parent=styles['Normal'],
                                   fontSize=10, fontName='Helvetica-Bold',
                                   alignment=TA_RIGHT, leading=12)
            g_val = ParagraphStyle('GV', parent=styles['Normal'],
                                   fontSize=10, fontName='Helvetica-Bold',
                                   alignment=TA_RIGHT, leading=12)

            logo_path = data['logo_path']
            logo_img = None
            if logo_path and os.path.exists(logo_path):
                try:
                    logo_img = Image(logo_path, width=1.05 * inch,
                                     height=1.05 * inch)
                except Exception:
                    logo_img = None

            addr_parts = self._split_address(data['address'])
            lines = [Paragraph(data['company_name'].upper(),
                               company_name_style)]
            if data['phone'] and data['email']:
                lines.append(Paragraph(
                    f"Phone: {data['phone']} | Email: {data['email']}",
                    company_info_style))
            lines.append(Paragraph(
                f"GST: {data['gst_no']} | PAN: {data['pan_no']} | "
                f"TAN: {data['tan_no']}",
                company_info_style))
            for part in addr_parts:
                lines.append(Paragraph(part, company_info_style))

            if logo_img:
                ht = Table([[logo_img, lines]],
                           colWidths=[2.6 * cm, 16.4 * cm])
                ht.setStyle(TableStyle([
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 0),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                ]))
                elements.append(ht)
            else:
                for line in lines:
                    elements.append(line)

            div = Table([[""]], colWidths=[19 * cm], rowHeights=[1])
            div.setStyle(TableStyle([
                ('LINEBELOW', (0, 0), (-1, -1), 1.2,
                 rl_colors.HexColor('#1a252f')),
            ]))
            elements.append(div)
            elements.append(Spacer(1, 4))

            elements.append(Paragraph("TAX INVOICE", title_style))
            elements.append(Spacer(1, 2))

            meta = Table(
                [[Paragraph(f"Invoice No: {data['invoice_no']}", meta_l),
                  Paragraph(
                      f"Date: {self._fmt_date_ddmmyyyy(data['issue_date'])}",
                      meta_r)]],
                colWidths=[9.5 * cm, 9.5 * cm])
            meta.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ]))
            elements.append(meta)
            elements.append(Spacer(1, 4))

            trav_parts = self._split_address(data['traveler_address'])
            bill_body = [Paragraph(data['traveler_name'], body_l)]
            for p in trav_parts:
                bill_body.append(Paragraph(p, body_l))
            if data['traveler_mobile']:
                bill_body.append(Paragraph(data['traveler_mobile'], body_l))
            if data['ref_no']:
                bill_body.append(Paragraph(f"Ref: {data['ref_no']}", body_l))

            pay_body = [Paragraph(data['company_name'], body_r)]
            for p in addr_parts:
                pay_body.append(Paragraph(p, body_r))
            if data['phone']:
                pay_body.append(Paragraph(data['phone'], body_r))
            if data['email']:
                pay_body.append(Paragraph(data['email'], body_r))

            bill_box = Table(
                [[Paragraph("BILL TO", box_lbl)], [bill_body]],
                colWidths=[9.2 * cm])
            bill_box.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, 0),
                 rl_colors.HexColor('#f0f0f0')),
                ('BOX', (0, 0), (-1, -1), 0.6,
                 rl_colors.HexColor('#bdc3c7')),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]))

            pay_box = Table(
                [[Paragraph("PAY TO", box_lbl)], [pay_body]],
                colWidths=[9.2 * cm])
            pay_box.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, 0),
                 rl_colors.HexColor('#f0f0f0')),
                ('BOX', (0, 0), (-1, -1), 0.6,
                 rl_colors.HexColor('#bdc3c7')),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]))

            addr_tbl = Table([[bill_box, "", pay_box]],
                             colWidths=[9.2 * cm, 0.6 * cm, 9.2 * cm])
            addr_tbl.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ]))
            elements.append(addr_tbl)
            elements.append(Spacer(1, 6))

            item_header = [
                Paragraph("S.No", hdr), Paragraph("Item", hdr),
                Paragraph("HSN", hdr), Paragraph("Rate", hdr),
                Paragraph("Qty", hdr), Paragraph("Amount", hdr),
                Paragraph("Disc%", hdr), Paragraph("Disc Amt", hdr),
                Paragraph("Taxable", hdr), Paragraph("CGST", hdr),
                Paragraph("SGST", hdr), Paragraph("Total", hdr),
            ]
            item_row = [
                Paragraph("1", cell_c),
                Paragraph(data['item_desc'], cell),
                Paragraph(data['hsn_code'], cell_c),
                Paragraph(f"{data['rate']:,.2f}", cell_r),
                Paragraph("1", cell_c),
                Paragraph(f"{data['base_amount']:,.2f}", cell_r),
                Paragraph(f"{data['discount_percent']:.2f}%", cell_c),
                Paragraph(f"{data['discount_amount']:,.2f}", cell_r),
                Paragraph(f"{data['taxable_value']:,.2f}", cell_r),
                Paragraph(f"{data['cgst']:,.2f}", cell_r),
                Paragraph(f"{data['sgst']:,.2f}", cell_r),
                Paragraph(f"{data['rounded_total']:,.2f}", cell_r),
            ]
            items = Table([item_header, item_row],
                          colWidths=[0.9 * cm, 4.7 * cm, 1.2 * cm, 1.5 * cm,
                                     0.8 * cm, 1.7 * cm, 1.1 * cm, 1.6 * cm,
                                     1.7 * cm, 1.5 * cm, 1.5 * cm, 1.8 * cm])
            items.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0),
                 rl_colors.HexColor('#2c3e50')),
                ('BACKGROUND', (0, 1), (-1, 1),
                 rl_colors.HexColor('#f8f9fa')),
                ('GRID', (0, 0), (-1, -1), 0.4,
                 rl_colors.HexColor('#bdc3c7')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ]))
            elements.append(items)
            elements.append(Spacer(1, 8))

            totals = [
                [Paragraph("Package Cost", t_lbl),
                 Paragraph(f"{data['base_amount']:,.2f}", t_val)],
                [Paragraph(f"Discount @ {data['discount_percent']:.2f}%",
                           t_disc_l),
                 Paragraph(f"- {data['discount_amount']:,.2f}", t_disc_v)],
                [Paragraph("Taxable Value", t_lbl),
                 Paragraph(f"{data['taxable_value']:,.2f}", t_val)],
                [Paragraph("GST", t_lbl),
                 Paragraph(f"{data['gst_amount']:,.2f}", t_val)],
                [Paragraph("Total Before TCS", t_lbl),
                 Paragraph(
                     f"{data['taxable_value'] + data['gst_amount']:,.2f}",
                     t_val)],
                [Paragraph("TCS", t_lbl),
                 Paragraph(f"{data['tcs_amount']:,.2f}", t_val)],
                [Paragraph("GRAND TOTAL", g_lbl),
                 Paragraph(f"{data['rounded_total']:,.0f}", g_val)],
            ]
            totals_inner = Table(totals, colWidths=[4.4 * cm, 3.4 * cm])
            totals_inner.setStyle(TableStyle([
                ('GRID', (0, 0), (-1, -1), 0.4,
                 rl_colors.HexColor('#bdc3c7')),
                ('BACKGROUND', (0, 1), (-1, 1),
                 rl_colors.HexColor('#fce4ec')),
                ('BACKGROUND', (0, 6), (-1, 6),
                 rl_colors.HexColor('#f9e79f')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 5),
                ('RIGHTPADDING', (0, 0), (-1, -1), 5),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ]))
            totals_outer = Table([["", totals_inner]],
                                 colWidths=[11.2 * cm, 7.8 * cm])
            totals_outer.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ]))
            elements.append(totals_outer)
            elements.append(Spacer(1, 6))

            words = number_to_words_indian(int(data['rounded_total']))
            if words.endswith(' Only'):
                words = words[:-5]
            ws = ParagraphStyle('W', parent=styles['Normal'], fontSize=8,
                                alignment=TA_RIGHT,
                                fontName='Helvetica-Oblique',
                                textColor=rl_colors.HexColor('#555'),
                                leading=10)
            wt = Table([["", Paragraph(
                f"Amount in Words: Rupees {words} Only", ws)]],
                colWidths=[11.2 * cm, 7.8 * cm])
            wt.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ]))
            elements.append(wt)
            elements.append(Spacer(1, 10))

            qr_box = Table(
                [[Paragraph("PAYMENT QR CODE", box_lbl)]],
                colWidths=[5.5 * cm], rowHeights=[2.6 * cm])
            qr_box.setStyle(TableStyle([
                ('BOX', (0, 0), (-1, -1), 0.6,
                 rl_colors.HexColor('#bdc3c7')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ]))

            bank_lines = [
                "<b>BANK DETAILS</b>",
                f"Bank Name : {data['bank_name']}",
                f"Account Holder : {data['company_name']}",
                f"Account Number : {data['account_no']}",
                f"IFSC Code : {data['ifsc']}",
                f"UPI ID : {data['upi']}",
            ]
            bank_body = "<br/>".join([l for l in bank_lines if l])
            bs = ParagraphStyle('BS', parent=styles['Normal'], fontSize=8,
                                alignment=TA_LEFT, leading=11)
            bank_box = Table([[Paragraph(bank_body, bs)]],
                             colWidths=[12.9 * cm])
            bank_box.setStyle(TableStyle([
                ('BOX', (0, 0), (-1, -1), 0.6,
                 rl_colors.HexColor('#bdc3c7')),
                ('LEFTPADDING', (0, 0), (-1, -1), 6),
                ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]))

            qb = Table([[qr_box, "", bank_box]],
                       colWidths=[5.5 * cm, 0.6 * cm, 12.9 * cm])
            qb.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ]))
            elements.append(qb)
            elements.append(Spacer(1, 8))

            decl = ParagraphStyle('D', parent=styles['Normal'], fontSize=7.5,
                                  alignment=TA_CENTER,
                                  fontName='Helvetica-Oblique',
                                  textColor=rl_colors.HexColor('#666'),
                                  leading=10)
            elements.append(Paragraph(
                "Declaration: TCS is collected as per Section 206C(1G) "
                "of the Income Tax Act, 1961. All disputes are subject "
                "to Chennai jurisdiction.", decl))
            elements.append(Spacer(1, 16))

            sign = ParagraphStyle('S', parent=styles['Normal'], fontSize=9,
                                  alignment=TA_RIGHT,
                                  fontName='Helvetica-Bold',
                                  leading=12)
            sign_sub = ParagraphStyle('SS', parent=styles['Normal'],
                                      fontSize=9, alignment=TA_RIGHT,
                                      leading=12)
            st = Table(
                [[Paragraph(f"For {data['company_name']}", sign)],
                 [Spacer(1, 30)],
                 [Paragraph("Authorised Signatory", sign_sub)]],
                colWidths=[19 * cm])
            st.setStyle(TableStyle([
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                ('TOPPADDING', (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]))
            elements.append(st)

            doc.build(elements)

            url = send_file_to_user(self.page, str(filepath), "Invoice PDF")
            self._snack(f"✅ PDF saved: {filepath.name}")
            if url:
                try:
                    self.page.launch_url(url)
                except Exception:
                    pass
        except Exception as ex:
            import traceback
            traceback.print_exc()
            self._snack(f"❌ PDF error: {ex}")

    def _split_address(self, address_text):
        if not address_text:
            return []
        parts = []
        for line in str(address_text).split('\n'):
            for p in line.split(','):
                if p.strip():
                    parts.append(p.strip())
        return parts

    def _get_invoice_data(self, invoice):
        company = (self.db.company_settings.iloc[0].to_dict()
                   if not self.db.company_settings.empty else {})
        bank = company.get('bank_details', {})
        if isinstance(bank, str):
            try:
                bank = json.loads(bank)
            except Exception:
                bank = {}

        tid = invoice.get('traveler_id', '')
        traveler = self.db.get_traveler_by_id(tid) if tid else None
        if traveler:
            tn = (f"{traveler.get('first_name', '')} "
                  f"{traveler.get('last_name', '')}").strip()
            ta = (traveler.get('mailing_address', '')
                  or traveler.get('passport_address', ''))
            tm = traveler.get('mobile', '')
            te = traveler.get('email', '')
            rn = traveler.get('file_reference', '')
        else:
            tn = ta = tm = te = rn = ''

        batch = self.db.get_batch_by_id(
            invoice.get('batch_id', '')) if invoice.get('batch_id') else None
        item_desc = (batch.get('batch_name', 'Package')
                     if batch else 'Package')

        base = self.safe_float(invoice.get('amount', 0))
        dp = self.safe_float(invoice.get('discount_percentage', 0))
        da = self.safe_float(invoice.get('discount_amount', 0))
        tv = self.safe_float(invoice.get('taxable_value', base - da))
        gp = self.safe_float(invoice.get('gst_percentage', 18))
        ga = self.safe_float(invoice.get('gst_amount', 0))
        tp = self.safe_float(invoice.get('tcs_percentage', 0.1))
        ta_ = self.safe_float(invoice.get('tcs_amount', 0))
        tot = self.safe_float(invoice.get('total_amount', 0))
        rt = self.safe_float(invoice.get('rounded_total', tot))

        return {
            'company_name': company.get('company_name', 'Alhudha Haj Travel'),
            'address': company.get('address', ''),
            'phone': company.get('phone', ''),
            'email': company.get('email', ''),
            'gst_no': company.get('gst_no', ''),
            'tan_no': company.get('tan_no', ''),
            'pan_no': company.get('pan_no', ''),
            'logo_path': company.get('logo_path', ''),
            'bank_name': bank.get('bank_name', ''),
            'account_no': bank.get('account_no', ''),
            'ifsc': bank.get('ifsc', '') or bank.get('swift', ''),
            'upi': bank.get('upi', '') or bank.get('iban', ''),
            'invoice_no': invoice.get('invoice_no', ''),
            'issue_date': invoice.get('issue_date', ''),
            'traveler_name': tn,
            'traveler_address': ta,
            'traveler_mobile': tm,
            'traveler_email': te,
            'ref_no': rn,
            'item_desc': item_desc,
            'hsn_code': '998555',
            'rate': base,
            'base_amount': base,
            'discount_percent': dp,
            'discount_amount': da,
            'taxable_value': tv,
            'gst_percentage': gp,
            'gst_amount': ga,
            'cgst': ga / 2 if ga else 0,
            'sgst': ga / 2 if ga else 0,
            'tcs_percentage': tp,
            'tcs_amount': ta_,
            'total_amount': tot,
            'rounded_total': rt,
        }

    # =============================================================================
    # 14.1.9 — Export Excel
    # =============================================================================
    def export_invoice_to_excel(self, invoice):
        try:
            import openpyxl
            from openpyxl.utils import get_column_letter
            from openpyxl.styles import (
                Font, Alignment, PatternFill, Border, Side)
            from openpyxl.drawing.image import Image as XLImage
            from PIL import Image as PILImage

            data = self._get_invoice_data(invoice)

            base = get_app_base_path()
            out_dir = Path(base) / "invoices"
            out_dir.mkdir(exist_ok=True)
            fname = (f"invoice_{data['invoice_no']}_"
                     f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
            filepath = out_dir / fname

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Invoice"

            logo_path = data.get('logo_path', '')
            if logo_path and os.path.exists(logo_path):
                try:
                    pil = PILImage.open(logo_path)
                    pil.thumbnail((90, 90))
                    buf = io.BytesIO()
                    pil.save(buf, format='PNG')
                    buf.seek(0)
                    xl = XLImage(buf)
                    xl.width = 90
                    xl.height = 90
                    ws.add_image(xl, 'A1')
                except Exception:
                    pass

            title_f = Font(name='Arial', size=14, bold=True, color='1A1A1A')
            sub_f = Font(name='Arial', size=9, color='555555')
            hdr_f = Font(name='Arial', size=10, bold=True, color='1A1A1A')
            nrm_f = Font(name='Arial', size=10, color='1A1A1A')
            bold_f = Font(name='Arial', size=10, bold=True, color='1A1A1A')
            disc_f = Font(name='Arial', size=10, bold=True, color='C2185B')
            grand_f = Font(name='Arial', size=12, bold=True, color='1A1A1A')

            r_align = Alignment(horizontal="right", vertical="center")
            l_align = Alignment(horizontal="left", vertical="center")
            c_align = Alignment(horizontal="center", vertical="center")

            thin = Side(style='thin', color='AAAAAA')
            borders = Border(left=thin, right=thin, top=thin, bottom=thin)
            disc_fill = PatternFill(start_color='FCE4EC', fill_type='solid')
            grand_fill = PatternFill(start_color='F9E79F', fill_type='solid')
            hdr_fill = PatternFill(start_color='E5E7E9', fill_type='solid')

            col_widths = {'A': 5, 'B': 30, 'C': 10, 'D': 14, 'E': 6,
                          'F': 15, 'G': 10, 'H': 14, 'I': 15, 'J': 13,
                          'K': 13, 'L': 15}
            for c, w in col_widths.items():
                ws.column_dimensions[c].width = w

            row = 1
            ws.merge_cells(f'B{row}:L{row}')
            ws[f'B{row}'] = data['company_name'].upper()
            ws[f'B{row}'].font = title_f
            ws[f'B{row}'].alignment = r_align
            row += 1

            ws.merge_cells(f'B{row}:L{row}')
            ws[f'B{row}'] = data['address']
            ws[f'B{row}'].font = sub_f
            ws[f'B{row}'].alignment = r_align
            row += 1

            ws.merge_cells(f'B{row}:L{row}')
            ws[f'B{row}'] = (f"Phone: {data['phone']} | "
                             f"Email: {data['email']}")
            ws[f'B{row}'].font = sub_f
            ws[f'B{row}'].alignment = r_align
            row += 1

            ws.merge_cells(f'B{row}:L{row}')
            ws[f'B{row}'] = (f"GSTIN: {data['gst_no']} | "
                             f"PAN: {data['pan_no']} | TAN: {data['tan_no']}")
            ws[f'B{row}'].font = sub_f
            ws[f'B{row}'].alignment = r_align
            row = 6

            ws.merge_cells(f'A{row}:L{row}')
            ws[f'A{row}'] = "TAX INVOICE"
            ws[f'A{row}'].font = Font(name='Arial', size=12, bold=True)
            ws[f'A{row}'].alignment = c_align
            row += 2

            ws[f'A{row}'] = f"Invoice No: {data['invoice_no']}"
            ws[f'A{row}'].font = hdr_f
            ws[f'A{row}'].alignment = l_align
            ws.merge_cells(f'A{row}:E{row}')

            ws.merge_cells(f'H{row}:L{row}')
            ws[f'H{row}'] = (f"Date: "
                             f"{self._fmt_date_ddmmyyyy(data['issue_date'])}")
            ws[f'H{row}'].font = hdr_f
            ws[f'H{row}'].alignment = r_align
            row += 2

            ws[f'A{row}'] = "BILL TO"
            ws[f'A{row}'].font = bold_f
            row += 1
            ws.merge_cells(f'A{row}:E{row}')
            ws[f'A{row}'] = data['traveler_name']
            ws[f'A{row}'].font = bold_f
            row += 1
            ws.merge_cells(f'A{row}:E{row}')
            ws[f'A{row}'] = data['traveler_address']
            ws[f'A{row}'].font = nrm_f
            row += 2

            hdr_row = row
            hdrs = ['S.No', 'Item', 'HSN', 'Rate', 'Qty', 'Amount',
                    'Disc%', 'Disc Amt', 'Taxable', 'CGST', 'SGST', 'Total']
            for i, h in enumerate(hdrs, 1):
                col = get_column_letter(i)
                c = ws[f'{col}{hdr_row}']
                c.value = h
                c.font = hdr_f
                c.fill = hdr_fill
                c.alignment = c_align
                c.border = borders
            row += 1

            vals = [
                1, data['item_desc'], data['hsn_code'], data['rate'],
                1, data['base_amount'], f"{data['discount_percent']:.2f}%",
                data['discount_amount'], data['taxable_value'],
                data['cgst'], data['sgst'], data['rounded_total']]
            for i, v in enumerate(vals, 1):
                col = get_column_letter(i)
                c = ws[f'{col}{row}']
                c.value = v
                c.font = disc_f if i in (7, 8) else nrm_f
                c.border = borders
                if i in (7, 8):
                    c.fill = disc_fill
                if i in (6, 8, 9, 12):
                    c.number_format = '₹ #,##0.00'
            row += 2

            totals = [
                ("Package Cost", data['base_amount'], False),
                (f"Discount @ {data['discount_percent']:.2f}%",
                 -data['discount_amount'], True),
                ("Taxable Value", data['taxable_value'], False),
                ("GST", data['gst_amount'], False),
                ("Total Before TCS",
                 data['taxable_value'] + data['gst_amount'], False),
                ("TCS", data['tcs_amount'], False),
                ("GRAND TOTAL", data['rounded_total'], "grand"),
            ]
            for label, val, flag in totals:
                ws.merge_cells(f'H{row}:K{row}')
                ws[f'H{row}'] = label
                ws[f'H{row}'].alignment = r_align
                ws[f'L{row}'] = val
                ws[f'L{row}'].number_format = '₹ #,##0.00'
                ws[f'L{row}'].alignment = r_align
                if flag is True:
                    ws[f'H{row}'].font = disc_f
                    ws[f'L{row}'].font = disc_f
                    ws[f'H{row}'].fill = disc_fill
                    ws[f'L{row}'].fill = disc_fill
                elif flag == "grand":
                    ws[f'H{row}'].font = grand_f
                    ws[f'L{row}'].font = grand_f
                    ws[f'H{row}'].fill = grand_fill
                    ws[f'L{row}'].fill = grand_fill
                else:
                    ws[f'H{row}'].font = bold_f
                row += 1

            wb.save(str(filepath))
            url = send_file_to_user(
                self.page, str(filepath), "Invoice Excel")
            self._snack(f"✅ Excel saved: {filepath.name}")
            if url:
                try:
                    self.page.launch_url(url)
                except Exception:
                    pass
        except Exception as ex:
            import traceback
            traceback.print_exc()
            self._snack(f"❌ Excel error: {ex}")

    def print_invoice(self, invoice):
        self._snack("🖨️ Generating PDF — press Ctrl+P when it opens")
        self.export_invoice_to_pdf(invoice)

    def _snack(self, msg):
        try:
            self.page.show_dialog(ft.SnackBar(content=ft.Text(msg)))
        except Exception as ex:
            print(f"[INVOICES] snack failed: {ex}")


# =================================================================================
# 14.2 — CLASS: InvoiceDialog
# =================================================================================
class InvoiceDialog:

    def __init__(self, page, db, current_user, on_save=None):
        self.page = page
        self.db = db
        self.current_user = current_user
        self.on_save = on_save
        self.settings_manager = SettingsManager(db)

        self.traveler_dd = None
        self.traveler_info = None
        self.batch_label = None
        self.payment_history = None
        self.amount_field = None
        self.discount_type_dd = None
        self.discount_value = None
        self.gst_dd = None
        self.tcs_field = None
        self.discount_amt_lbl = None
        self.discount_pct_lbl = None
        self.taxable_lbl = None
        self.total_lbl = None
        self.rounded_lbl = None
        self.breakdown_lbl = None
        self.notes_field = None
        self.generate_btn = None
        self.dialog = None
        self.traveler_map = {}
        self._last_applied_tid = None
        self._selection_watch_running = False

        self._computed_discount_pct = 0.0
        self._computed_discount_amount = 0.0
        self._computed_taxable_value = 0.0
        self._typing_until = 0.0

        self.setup_ui()
        self.load_tax_rates()

    def _field(self, **kwargs):
        if "content_padding" not in kwargs:
            kwargs["content_padding"] = ft.Padding.symmetric(
                horizontal=10, vertical=10)
        field = ft.TextField(**kwargs)

        def _on_change(e, _f=field):
            self._typing_until = time.time() + TYPING_GUARD_SECONDS
            try:
                ctrl = getattr(e, "control", None)
                if ctrl is not None and hasattr(ctrl, "value"):
                    if _f.value != ctrl.value:
                        _f.value = ctrl.value
            except Exception:
                pass

        field.on_change = _on_change
        return field

    def setup_ui(self):
        traveler_opts = [ft.dropdown.Option(key="", text="Select Traveler")]
        self.traveler_map = {}
        try:
            for i, t in enumerate(self.db.get_travelers()):
                if not t.get('id'):
                    continue
                name = (f"{t.get('first_name', '')} "
                        f"{t.get('last_name', '')}").strip()
                pp = t.get('passport_no', '')
                label = f"{name} ({pp})" if pp else name
                key = f"t_{i}"
                self.traveler_map[key] = {
                    'id': str(t['id']).strip(),
                    'text': label
                }
                traveler_opts.append(
                    ft.dropdown.Option(key=key, text=label))
        except Exception:
            pass

        self.traveler_dd = ft.Dropdown(
            label="👤 Traveler *",
            options=traveler_opts, value="",
            height=48, text_size=12)
        self.traveler_dd.on_change = self.on_traveler_change

        self.traveler_info = ft.Text(
            "No traveler selected", size=11, color=ft.Colors.GREY_600)
        self.batch_label = ft.Text(
            "Not assigned", size=12,
            weight=ft.FontWeight.BOLD, color="#2c3e50")
        self.payment_history = ft.Text(
            "💳 Payment history: select a traveler to view payments.",
            size=11, color=ft.Colors.GREY_700)

        self.amount_field = self._field(
            label="💰 Base Amount (₹) *", value="0",
            height=48, text_size=12)

        def _amount_on_change(e, _f=self.amount_field):
            self._typing_until = time.time() + TYPING_GUARD_SECONDS
            try:
                ctrl = getattr(e, "control", None)
                if ctrl is not None and hasattr(ctrl, "value"):
                    if _f.value != ctrl.value:
                        _f.value = ctrl.value
            except Exception:
                pass
            self.recalculate(e)

        self.amount_field.on_change = _amount_on_change

        self.discount_type_dd = ft.Dropdown(
            label="🎁 Discount Type",
            options=[
                ft.dropdown.Option(key="pct", text="Percentage (%)"),
                ft.dropdown.Option(key="amt", text="Fixed Amount (₹)"),
            ],
            value="pct", height=48, text_size=12)
        self.discount_type_dd.on_change = self.on_disc_type_change

        self.discount_value = self._field(
            label="🎁 Discount Value", value="0",
            height=48, text_size=12)

        def _discval_on_change(e, _f=self.discount_value):
            self._typing_until = time.time() + TYPING_GUARD_SECONDS
            try:
                ctrl = getattr(e, "control", None)
                if ctrl is not None and hasattr(ctrl, "value"):
                    if _f.value != ctrl.value:
                        _f.value = ctrl.value
            except Exception:
                pass
            self.recalculate(e)

        self.discount_value.on_change = _discval_on_change

        self.gst_dd = ft.Dropdown(
            label="📊 GST %",
            options=[ft.dropdown.Option(x) for x in
                     ["0", "5", "12", "18", "28"]],
            value="18", height=48, text_size=12)
        self.gst_dd.on_change = self.recalculate

        self.tcs_field = self._field(
            label="📊 TCS %", value="0.1",
            height=48, text_size=12)

        def _tcs_on_change(e, _f=self.tcs_field):
            self._typing_until = time.time() + TYPING_GUARD_SECONDS
            try:
                ctrl = getattr(e, "control", None)
                if ctrl is not None and hasattr(ctrl, "value"):
                    if _f.value != ctrl.value:
                        _f.value = ctrl.value
            except Exception:
                pass
            self.recalculate(e)

        self.tcs_field.on_change = _tcs_on_change

        self.discount_amt_lbl = ft.Text(
            "Discount Amount: ₹0.00", size=12,
            weight=ft.FontWeight.BOLD, color="#c2185b")
        self.discount_pct_lbl = ft.Text(
            "Effective Discount %: 0.00%", size=11, color="#c2185b")
        self.taxable_lbl = ft.Text(
            "Taxable Value: ₹0.00", size=12,
            weight=ft.FontWeight.BOLD, color="#2c3e50")
        self.total_lbl = ft.Text(
            "Total: ₹0.00", size=13, weight=ft.FontWeight.BOLD)
        self.rounded_lbl = ft.Text(
            "★ Rounded Total: ₹0", size=15,
            weight=ft.FontWeight.BOLD, color="#e67e22")
        self.breakdown_lbl = ft.Text(
            "GST: ₹0.00 | TCS: ₹0.00", size=11, color=ft.Colors.GREY_600)

        self.notes_field = self._field(
            label="📝 Notes", multiline=True,
            min_lines=2, max_lines=3, text_size=12)

        self.generate_btn = ft.Button(
            content=ft.Text("💾 Generate Invoice"),
            on_click=self.save, bgcolor="#27ae60",
            color=ft.Colors.WHITE, height=45, disabled=True)

        content = ft.Column(
            controls=[
                self.traveler_dd,
                self.traveler_info,
                self.batch_label,
                ft.Container(
                    content=self.payment_history,
                    bgcolor="#f5f7fa", padding=8, border_radius=6),
                self.amount_field,
                self.discount_type_dd,
                self.discount_value,
                self.gst_dd,
                self.tcs_field,
                ft.Divider(height=10),
                self.discount_amt_lbl,
                self.discount_pct_lbl,
                self.taxable_lbl,
                self.total_lbl,
                self.rounded_lbl,
                self.breakdown_lbl,
                self.notes_field,
            ],
            spacing=10,
            scroll=ft.ScrollMode.AUTO,
        )

        self.dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("📄 Generate Invoice",
                          weight=ft.FontWeight.BOLD, size=15),
            content=ft.Container(content=content,
                                 width=560, height=680, padding=10),
            actions=[
                ft.TextButton(content=ft.Text("Cancel"),
                              on_click=lambda e: self.page.pop_dialog()),
                self.generate_btn,
            ],
            actions_alignment=ft.MainAxisAlignment.END,
        )

    def load_tax_rates(self):
        tax = self.settings_manager.get_tax_settings()
        try:
            gst = _rate(tax.get('gst_percentage', 18))
            tcs = _rate(tax.get('tcs_percentage', 0.1))
        except Exception:
            gst = Decimal('18')
            tcs = Decimal('0.1')
        self.gst_dd.value = (
            str(int(gst)) if int(gst) in [0, 5, 12, 18, 28] else "18")
        self.tcs_field.value = str(tcs)
        self.recalculate(None)

    def on_traveler_change(self, e):
        # Apply the event value immediately. Some Flet web versions update
        # the visible label before traveler_dd.value is synchronized.
        control = getattr(e, 'control', None)
        event_val = (getattr(control, 'value', None)
                     or getattr(e, 'data', None)
                     or '')
        print(f"[InvoiceDialog] traveler dropdown event: value={event_val!r}, "
              f"control_value={getattr(control, 'value', None)!r}", flush=True)
        immediate_tid = self._resolve_tid(event_val)
        if immediate_tid:
            self._last_applied_tid = immediate_tid
            self._apply_traveler_change(immediate_tid)

        async def _process():
            # Keep polling for delayed browser-side Dropdown synchronization.
            for _ in range(15):
                await asyncio.sleep(0.1)
                values = [
                    getattr(self.traveler_dd, 'value', None),
                    getattr(control, 'value', None),
                    getattr(e, 'data', None),
                    event_val,
                ]
                for val in values:
                    tid = self._resolve_tid(val)
                    if tid and tid != getattr(
                            self, '_last_applied_tid', None):
                        self._last_applied_tid = tid
                        self._apply_traveler_change(tid)
                        return
                if any(values):
                    return
            print(f"[DEBUG-INV] FAILED to resolve traveler. val={event_val!r}")

        try:
            self.page.run_task(_process)
        except Exception as ex:
            # Immediate resolution remains usable on older Flet runtimes.
            print(f"[DEBUG-INV] polling task failed: {ex}")

    def _resolve_tid(self, val):
        if not val:
            return None
        val_str = str(val).strip()
        
        # 1. Check map by key
        if val_str in self.traveler_map:
            return self.traveler_map[val_str]['id']
            
        # 2. Check map by exact text
        for k, info in self.traveler_map.items():
            if info['text'] == val_str:
                return info['id']
                
        # 3. Check database directly (in case of formatting discrepancies)
        for t in self.db.get_travelers():
            name = f"{t.get('first_name', '')} {t.get('last_name', '')}".strip()
            pp = t.get('passport_no', '')
            label = f"{name} ({pp})" if pp else name
            if label == val_str or name == val_str or str(t['id']) == val_str:
                return str(t['id']).strip()
                
            # Fuzzy match
            if val_str in label or val_str in name:
                 return str(t['id']).strip()
                 
        return None

    def _apply_traveler_change(self, tid):
        if not tid:
            self.traveler_info.value = "No traveler selected"
            self.traveler_info.color = ft.Colors.GREY_600
            self.batch_label.value = "Not assigned"
            self.amount_field.value = "0"
            self.payment_history.value = (
                "💳 Payment history: select a traveler to view payments.")
            self.generate_btn.disabled = True
            self.recalculate(None)
            self._force_ui_update()
            return

        t = self.db.get_traveler_by_id(tid)
        if not t:
            self.traveler_info.value = "⚠️ Traveler not found"
            self.traveler_info.color = ft.Colors.RED_600
            self.batch_label.value = "Batch not found"
            self.amount_field.value = "0"
            self.payment_history.value = "💳 Payment history unavailable."
            self.generate_btn.disabled = True
            self.recalculate(None)
            self._force_ui_update()
            return

        existing_invoices = self.db.get_invoices(traveler_id=tid)
        self._update_payment_history(tid)
        name = (f"{t.get('first_name', '')} "
                f"{t.get('last_name', '')}").strip()

        if existing_invoices:
            inv_no = existing_invoices[0].get('invoice_no', '')
            self.traveler_info.value = (
                f"⚠️ {name} already has invoice {inv_no}.")
            self.traveler_info.color = ft.Colors.RED_600
            self.batch_label.value = "Cannot create duplicate"
            self.amount_field.value = "0"
            self.generate_btn.disabled = True
        else:
            self.traveler_info.value = (
                f"✅ {name} | Passport: {t.get('passport_no', '')}")
            self.traveler_info.color = ft.Colors.GREY_600

            bid = str(t.get('batch_id') or '').strip()
            print(f"[InvoiceDialog] traveler batch ID: {bid!r}")
            batch = None
            if bid:
                try:
                    batch = self.db.get_batch_by_id(bid)
                except Exception as ex:
                    print(f"[InvoiceDialog] direct batch lookup failed: {ex}")
                if not batch:
                    batch = next(
                        (b for b in self.db.get_batches()
                         if str(b.get('id') or '').strip() == bid),
                        None)
            print(f"[InvoiceDialog] matched batch: {batch}")

            if batch:
                price = _money(batch.get('price', 0) or 0)
                self.batch_label.value = (
                    f"📦 {batch.get('batch_name', '')} | "
                    f"Price: ₹{price:,.2f}")
                self.amount_field.value = str(price)
                self.generate_btn.disabled = price <= 0
            elif bid:
                self.batch_label.value = f"Batch not found: {bid}"
                self.amount_field.value = "0"
                self.generate_btn.disabled = True
            else:
                self.batch_label.value = "No batch assigned"
                self.amount_field.value = "0"
                self.generate_btn.disabled = True

        self.recalculate(None)
        self._force_ui_update()

    def _force_ui_update(self):
        try:
            self.dialog.update()
        except Exception as ex:
            print(f"[DEBUG-INV] Dialog update error: {ex}")

    def _update_payment_history(self, traveler_id):
        """Display payment records without changing payment or receipt data."""
        try:
            payments = list(self.db.get_payments(traveler_id) or [])
        except Exception as ex:
            self.payment_history.value = f"💳 Payment history error: {ex}"
            return

        if not payments:
            self.payment_history.value = (
                "💳 Payment history: no payments recorded.")
            return

        payments.sort(
            key=lambda p: str(p.get('payment_date', '') or ''), reverse=True)
        total = Decimal('0.00')
        lines = [f"💳 Payment history ({len(payments)} record(s)):"]
        for payment in payments:
            amount = _money(payment.get('amount', 0) or 0)
            total += amount
            date = str(payment.get('payment_date', '') or '')[:10] or '—'
            status = str(payment.get('status', '') or '—')
            method = str(payment.get('payment_method', '') or '—')
            lines.append(
                f"• {date} | ₹{amount:,.2f} | {method} | {status}")
        lines.append(f"Total paid: ₹{total:,.2f}")
        self.payment_history.value = "\n".join(lines)


    async def _watch_traveler_selection(self):
        """Fallback for Flet web builds that update the dropdown without
        reliably firing the on_change callback."""
        if self._selection_watch_running:
            return
        self._selection_watch_running = True
        print("[DEBUG-INV] traveler watcher started", flush=True)
        try:
            for _ in range(240):
                await asyncio.sleep(0.25)
                value = getattr(self.traveler_dd, 'value', None)
                tid = self._resolve_tid(value)
                if tid and tid != self._last_applied_tid:
                    print(f"[DEBUG-INV] watcher selected traveler: {tid!r}")
                    self._apply_traveler_change(tid)
                    self._last_applied_tid = tid
        except Exception as ex:
            print(f"[DEBUG-INV] traveler watcher failed: {ex}")
        finally:
            self._selection_watch_running = False

    def on_disc_type_change(self, e):
        if self.discount_type_dd.value == "amt":
            self.discount_value.label = "🎁 Discount Amount (₹)"
        else:
            self.discount_value.label = "🎁 Discount Value (%)"
        self.discount_value.value = "0"
        self.recalculate(None)
        self._force_ui_update()

    def recalculate(self, e):
        try:
            amount = _money(self.amount_field.value or 0)
        except (ValueError, TypeError):
            amount = Decimal('0.00')
        try:
            dval = _rate(self.discount_value.value or 0)
        except (ValueError, TypeError):
            dval = Decimal('0.0000')
        try:
            gst = _rate(self.gst_dd.value or 0)
        except (ValueError, TypeError):
            gst = Decimal('0.0000')
        try:
            tcs = _rate(self.tcs_field.value or 0)
        except (ValueError, TypeError):
            tcs = Decimal('0.0000')

        if amount < 0: amount = Decimal('0.00')
        if gst < 0: gst = Decimal('0.0000')
        if tcs < 0: tcs = Decimal('0.0000')
        if dval < 0: dval = Decimal('0.0000')
        if gst > Decimal('100'): gst = Decimal('100')
        if tcs > Decimal('100'): tcs = Decimal('100')

        if self.discount_type_dd.value == "pct":
            disc_pct = dval
            disc_amt = amount * disc_pct / Decimal('100')
        else:
            disc_amt = dval
            disc_pct = (disc_amt / amount * Decimal('100')) if amount > 0 else Decimal('0')

        if disc_pct > Decimal('100'): disc_pct = Decimal('100')
        if disc_amt > amount: disc_amt = amount

        taxable = amount - disc_amt
        gst_amt = taxable * gst / Decimal('100')
        taxable_with_gst = taxable + gst_amt
        tcs_amt = taxable_with_gst * tcs / Decimal('100')
        total = taxable_with_gst + tcs_amt
        rounded = _round_rule(total)

        self._computed_discount_pct = disc_pct
        self._computed_discount_amount = disc_amt
        self._computed_taxable_value = taxable

        self.discount_amt_lbl.value = f"Discount Amount: ₹{disc_amt:,.2f}"
        self.discount_pct_lbl.value = (
            f"Effective Discount %: {disc_pct:.2f}%")
        self.taxable_lbl.value = f"Taxable Value: ₹{taxable:,.2f}"
        self.total_lbl.value = f"Total: ₹{total:,.2f}"
        self.rounded_lbl.value = f"★ Rounded Total: ₹{rounded:,.0f}"
        self.breakdown_lbl.value = (f"GST: ₹{gst_amt:,.2f} | "
                                    f"TCS: ₹{tcs_amt:,.2f}")
        
        for c in (self.discount_amt_lbl, self.discount_pct_lbl,
                  self.taxable_lbl, self.total_lbl, self.rounded_lbl,
                  self.breakdown_lbl):
            try:
                c.update()
            except Exception:
                pass

    def save(self, e):
        val = getattr(self.traveler_dd, 'value', None)
        tid = (getattr(self, '_last_applied_tid', None)
               or self._resolve_tid(val))
                    
        if not tid:
            self._snack("⚠️ Select a traveler")
            return
            
        traveler = self.db.get_traveler_by_id(tid)
        bid = traveler.get('batch_id') if traveler else None
        if not bid:
            self._snack("⚠️ Traveler has no batch")
            return
            
        existing_invoices = self.db.get_invoices(traveler_id=tid)
        if existing_invoices:
            self._snack(f"⚠️ Invoice {existing_invoices[0].get('invoice_no')} already exists for this traveler.")
            return

        try:
            amount = _money(self.amount_field.value or 0)
        except (ValueError, TypeError):
            amount = Decimal('0.00')
        if amount <= 0:
            self._snack("⚠️ Amount must be greater than zero")
            return
            
        try:
            gst = _rate(self.gst_dd.value or 18)
            tcs = _rate(self.tcs_field.value or 0.1)
            if gst < 0 or tcs < 0:
                self._snack("⚠️ GST/TCS cannot be negative")
                return
        except Exception:
            gst = Decimal('18')
            tcs = Decimal('0.1')

        is_pct = self.discount_type_dd.value == "pct"

        try:
            if is_pct:
                inv_id = self.db.generate_invoice(
                    tid, bid, amount, gst, tcs,
                    discount_percentage=self._computed_discount_pct,
                    discount_amount=None)
            else:
                inv_id = self.db.generate_invoice(
                    tid, bid, amount, gst, tcs,
                    discount_percentage=Decimal('0'),
                    discount_amount=self._computed_discount_amount)
            try:
                self.db.log_activity(
                    self.current_user['id'], "generate_invoice",
                    f"Generated invoice for {traveler.get('first_name', '')}")
            except Exception:
                pass
            self._typing_until = 0.0
            self.page.pop_dialog()
            self._snack(f"✅ Invoice {inv_id} generated")
            if self.on_save:
                self.on_save()
        except Exception as ex:
            import traceback
            traceback.print_exc()
            self._snack(f"❌ {ex}")

    def show(self):
        self.page.show_dialog(self.dialog)
        try:
            self.page.run_task(self._watch_traveler_selection)
        except Exception as ex:
            print(f"[DEBUG-INV] traveler watcher could not start: {ex}")

    def _snack(self, msg):
        try:
            self.page.show_dialog(ft.SnackBar(content=ft.Text(msg)))
        except Exception as ex:
            print(f"[INV-DIALOG] snack failed: {ex}")


# =================================================================================
# 14.3 — CLASS: InvoiceModifyDialog
# =================================================================================
class InvoiceModifyDialog:

    def __init__(self, page, db, current_user, invoice, on_save=None):
        self.page = page
        self.db = db
        self.current_user = current_user
        self.invoice = dict(invoice)
        self.on_save = on_save

        self.amount_field = None
        self.disc_type_dd = None
        self.disc_value = None
        self.gst_dd = None
        self.tcs_field = None
        self.status_dd = None
        self.notes_field = None
        self.disc_amt_lbl = None
        self.disc_pct_lbl = None
        self.taxable_lbl = None
        self.rounded_lbl = None
        self.dialog = None
        self._computed = {}
        self._typing_until = 0.0

        self.setup_ui()
        self.load_data()

    def _field(self, **kwargs):
        if "content_padding" not in kwargs:
            kwargs["content_padding"] = ft.Padding.symmetric(
                horizontal=10, vertical=10)
        field = ft.TextField(**kwargs)

        def _on_change(e, _f=field):
            self._typing_until = time.time() + TYPING_GUARD_SECONDS
            try:
                ctrl = getattr(e, "control", None)
                if ctrl is not None and hasattr(ctrl, "value"):
                    if _f.value != ctrl.value:
                        _f.value = ctrl.value
            except Exception:
                pass

        field.on_change = _on_change
        return field

    def setup_ui(self):
        is_paid = str(self.invoice.get('status', '')).lower() == 'paid'

        self.amount_field = self._field(
            label="Base Amount (₹)", height=48, text_size=12, disabled=is_paid)

        def _amt_change(e, _f=self.amount_field):
            self._typing_until = time.time() + TYPING_GUARD_SECONDS
            try:
                ctrl = getattr(e, "control", None)
                if ctrl is not None and hasattr(ctrl, "value"):
                    if _f.value != ctrl.value:
                        _f.value = ctrl.value
            except Exception:
                pass
            self.recalc(e)

        self.amount_field.on_change = _amt_change

        self.disc_type_dd = ft.Dropdown(
            label="Discount Type",
            options=[
                ft.dropdown.Option(key="pct", text="Percentage (%)"),
                ft.dropdown.Option(key="amt", text="Fixed Amount (₹)"),
            ], value="pct", height=48, text_size=12, disabled=is_paid)
        self.disc_type_dd.on_change = self.recalc

        self.disc_value = self._field(
            label="Discount Value", value="0",
            height=48, text_size=12, disabled=is_paid)

        def _dv_change(e, _f=self.disc_value):
            self._typing_until = time.time() + TYPING_GUARD_SECONDS
            try:
                ctrl = getattr(e, "control", None)
                if ctrl is not None and hasattr(ctrl, "value"):
                    if _f.value != ctrl.value:
                        _f.value = ctrl.value
            except Exception:
                pass
            self.recalc(e)

        self.disc_value.on_change = _dv_change

        self.gst_dd = ft.Dropdown(
            label="GST %",
            options=[ft.dropdown.Option(x) for x in
                     ["0", "5", "12", "18", "28"]],
            value="18", height=48, text_size=12, disabled=is_paid)
        self.gst_dd.on_change = self.recalc

        self.tcs_field = self._field(
            label="TCS %", value="0.1",
            height=48, text_size=12, disabled=is_paid)

        def _tcs_change(e, _f=self.tcs_field):
            self._typing_until = time.time() + TYPING_GUARD_SECONDS
            try:
                ctrl = getattr(e, "control", None)
                if ctrl is not None and hasattr(ctrl, "value"):
                    if _f.value != ctrl.value:
                        _f.value = ctrl.value
            except Exception:
                pass
            self.recalc(e)

        self.tcs_field.on_change = _tcs_change

        self.status_dd = ft.Dropdown(
            label="Status",
            options=[ft.dropdown.Option(x) for x in
                     ["pending", "paid", "overdue", "cancelled"]],
            value="pending", height=48, text_size=12)

        self.notes_field = self._field(
            label="Notes", multiline=True,
            min_lines=2, max_lines=3, text_size=12)

        self.disc_amt_lbl = ft.Text(
            "Discount Amount: ₹0.00", size=12,
            weight=ft.FontWeight.BOLD, color="#c2185b")
        self.disc_pct_lbl = ft.Text(
            "Effective Discount %: 0.00%", size=11, color="#c2185b")
        self.taxable_lbl = ft.Text(
            "Taxable Value: ₹0.00", size=12,
            weight=ft.FontWeight.BOLD)
        self.rounded_lbl = ft.Text(
            "★ Rounded Total: ₹0", size=15,
            weight=ft.FontWeight.BOLD, color="#e67e22")

        content = ft.Column(
            controls=[
                ft.Text(f"Invoice: {self.invoice.get('invoice_no', '')}",
                        size=13, weight=ft.FontWeight.BOLD),
                self.amount_field,
                self.disc_type_dd,
                self.disc_value,
                self.gst_dd,
                self.tcs_field,
                self.status_dd,
                ft.Divider(height=10),
                self.disc_amt_lbl,
                self.disc_pct_lbl,
                self.taxable_lbl,
                self.rounded_lbl,
                self.notes_field,
            ], spacing=10, scroll=ft.ScrollMode.AUTO)

        self.dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("✏️ Modify Invoice",
                          weight=ft.FontWeight.BOLD, size=15),
            content=ft.Container(content=content,
                                 width=520, height=620, padding=10),
            actions=[
                ft.TextButton(content=ft.Text("Cancel"),
                              on_click=lambda e: self.page.pop_dialog()),
                ft.Button(content=ft.Text("💾 Save Changes"),
                          on_click=self.save,
                          bgcolor="#27ae60", color=ft.Colors.WHITE),
            ],
            actions_alignment=ft.MainAxisAlignment.END,
        )

    def load_data(self):
        inv = self.invoice
        self.amount_field.value = f"{self._f(inv.get('amount', 0)):.2f}"
        existing_pct = self._f(inv.get('discount_percentage', 0))
        existing_amt = self._f(inv.get('discount_amount', 0))

        use_amt = (abs(existing_pct - round(existing_pct, 2)) > Decimal('0.001')
                   and existing_amt > 0)
        if use_amt:
            self.disc_type_dd.value = "amt"
            self.disc_value.value = f"{existing_amt:.2f}"
        else:
            self.disc_type_dd.value = "pct"
            self.disc_value.value = f"{existing_pct:.2f}"

        self.gst_dd.value = str(int(self._f(inv.get('gst_percentage', 18))))
        self.tcs_field.value = f"{self._f(inv.get('tcs_percentage', 0.1)):.2f}"
        self.status_dd.value = str(inv.get('status', 'pending'))
        self.notes_field.value = str(inv.get('notes', ''))
        self.recalc(None)

    def _f(self, v, d=Decimal('0')):
        try:
            return Decimal(str(v))
        except (InvalidOperation, ValueError, TypeError):
            return d

    def recalc(self, e=None):
        try:
            amount = _money(self.amount_field.value or 0)
        except (ValueError, TypeError):
            amount = Decimal('0.00')
        try:
            dval = _rate(self.disc_value.value or 0)
        except (ValueError, TypeError):
            dval = Decimal('0.0000')
        try:
            gst = _rate(self.gst_dd.value or 0)
        except (ValueError, TypeError):
            gst = Decimal('0.0000')
        try:
            tcs = _rate(self.tcs_field.value or 0)
        except (ValueError, TypeError):
            tcs = Decimal('0.0000')

        if amount < 0: amount = Decimal('0.00')
        if gst < 0: gst = Decimal('0.0000')
        if tcs < 0: tcs = Decimal('0.0000')
        if dval < 0: dval = Decimal('0.0000')

        if self.disc_type_dd.value == "pct":
            disc_pct = dval
            disc_amt = amount * disc_pct / Decimal('100')
        else:
            disc_amt = dval
            disc_pct = (disc_amt / amount * Decimal('100')) if amount > 0 else Decimal('0')

        if disc_pct > Decimal('100'): disc_pct = Decimal('100')
        if disc_amt > amount: disc_amt = amount

        taxable = amount - disc_amt
        gst_amt = taxable * gst / Decimal('100')
        taxable_with_gst = taxable + gst_amt
        tcs_amt = taxable_with_gst * tcs / Decimal('100')
        total = taxable_with_gst + tcs_amt
        rounded = _round_rule(total)

        self._computed = {
            'discount_pct': disc_pct, 'discount_amt': disc_amt,
            'taxable': taxable, 'gst': gst, 'gst_amt': gst_amt,
            'tcs': tcs, 'tcs_amt': tcs_amt,
            'total': total, 'rounded': rounded}

        self.disc_amt_lbl.value = f"Discount Amount: ₹{disc_amt:,.2f}"
        self.disc_pct_lbl.value = f"Effective Discount %: {disc_pct:.2f}%"
        self.taxable_lbl.value = f"Taxable Value: ₹{taxable:,.2f}"
        self.rounded_lbl.value = f"★ Rounded Total: ₹{rounded:,.0f}"
        
        for c in (self.disc_amt_lbl, self.disc_pct_lbl,
                  self.taxable_lbl, self.rounded_lbl):
            try:
                c.update()
            except Exception:
                pass

    def save(self, e):
        try:
            c = self._computed
            if not c:
                self.recalc(None)
                c = self._computed

            if c['taxable'] < 0 or c['total'] < 0:
                self._snack("⚠️ Invalid invoice totals")
                return
            
            new_status = (self.status_dd.value or 'pending').lower()
            old_status = str(self.invoice.get('status', 'pending')).lower()

            update = {
                'amount': _money(self.amount_field.value or 0),
                'discount_percentage': c['discount_pct'],
                'discount_amount': c['discount_amt'],
                'taxable_value': c['taxable'],
                'gst_percentage': c['gst'],
                'gst_amount': c['gst_amt'],
                'tcs_percentage': c['tcs'],
                'tcs_amount': c['tcs_amt'],
                'total_amount': c['total'],
                'rounded_total': c['rounded'],
                'status': new_status,
                'notes': (self.notes_field.value or '').strip(),
            }

            if hasattr(self.db, "update_invoice"):
                self.db.update_invoice(self.invoice['id'], **update)
            else:
                for k, v in update.items():
                    self.db.invoices.loc[
                        self.db.invoices['id'] == self.invoice['id'], k] = v
                self.db._save_df(self.db.invoices, "invoices.csv")

            # Invoice status edits do not create, delete, or alter payments.
            # Payments and receipts are managed only from their own tabs.

            try:
                self.db.log_activity(
                    self.current_user['id'], "modify_invoice",
                    f"Modified invoice {self.invoice.get('invoice_no', '')} "
                    f"({old_status} → {new_status})")
            except Exception:
                pass

            self._typing_until = 0.0
            self.page.pop_dialog()
            self._snack("✅ Invoice updated")
            if self.on_save:
                self.on_save()
        except Exception as ex:
            import traceback
            traceback.print_exc()
            self._snack(f"❌ {ex}")

    def _cascade_mark_paid(self, traveler_id, invoice_id, amount):
        if not traveler_id or amount is None or amount <= 0:
            return
        try:
            traveler = self.db.get_traveler_by_id(traveler_id)
            batch_id = traveler.get('batch_id', '') if traveler else ''

            existing = None
            try:
                payments = self.db.get_payments(traveler_id) or []
            except Exception:
                payments = []
            for p in payments:
                if str(p.get('invoice_id', '')) == str(invoice_id):
                    existing = p
                    break

            if existing:
                try:
                    cur = _money(existing.get('amount', 0) or 0)
                except Exception:
                    cur = Decimal('0.00')
                if abs(cur - _money(amount)) > Decimal('0.01'):
                    if hasattr(self.db, "update_payment"):
                        self.db.update_payment(
                            existing['id'],
                            {
                                "amount": _money(amount),
                                "notes": "Auto-synced via invoice status",
                                "status": "completed"
                            })
                    else:
                        self.db.payments.loc[
                            self.db.payments['id'] == existing['id'],
                            'amount'] = _money(amount)
                        self.db.payments.loc[
                            self.db.payments['id'] == existing['id'],
                            'status'] = 'completed'
                        self.db._save_df(self.db.payments, "payments.csv")
                else:
                    if hasattr(self.db, "update_payment"):
                        self.db.update_payment(existing['id'], {"status": "completed"})
                    else:
                        self.db.payments.loc[
                            self.db.payments['id'] == existing['id'],
                            'status'] = 'completed'
                        self.db._save_df(self.db.payments, "payments.csv")
            else:
                payload = {
                    'traveler_id': traveler_id,
                    'batch_id': batch_id,
                    'invoice_id': invoice_id,
                    'amount': _money(amount),
                    'payment_date': datetime.now().strftime("%Y-%m-%d"),
                    'payment_method': 'Auto (invoice status)',
                    'notes': 'Auto-created when invoice marked paid',
                    'status': 'completed'
                }
                if hasattr(self.db, "add_payment"):
                    try:
                        self.db.add_payment(payload)
                    except TypeError:
                        self.db.add_payment(**payload)
                else:
                    try:
                        self.db.payments = pd.concat(
                            [self.db.payments,
                             pd.DataFrame([payload])],
                            ignore_index=True)
                        self.db._save_df(self.db.payments, "payments.csv")
                    except Exception as _e:
                        print(f"[cascade_mark_paid fallback] {_e}")
        except Exception as ex:
            print(f"[cascade_mark_paid] {ex}")

    def _cascade_unmark_paid(self, traveler_id, invoice_id):
        if not traveler_id:
            return
        try:
            try:
                payments = self.db.get_payments(traveler_id) or []
            except Exception:
                payments = []

            target = None
            for p in payments:
                if (str(p.get('invoice_id', '')) == str(invoice_id)
                        and 'auto' in str(p.get('payment_method', '')).lower()):
                    target = p
                    break

            if not target:
                return

            if hasattr(self.db, "update_payment"):
                self.db.update_payment(
                    target['id'],
                    {
                        "status": "refunded",
                        "notes": "Auto-voided: Invoice marked unpaid"
                    }
                )
            else:
                self.db.payments.loc[
                    self.db.payments['id'] == target['id'],
                    'status'] = 'refunded'
                self.db._save_df(self.db.payments, "payments.csv")
                
        except Exception as ex:
            print(f"[cascade_unmark_paid] {ex}")

    def show(self):
        self.page.show_dialog(self.dialog)

    def _snack(self, msg):
        try:
            self.page.show_dialog(ft.SnackBar(content=ft.Text(msg)))
        except Exception:
            pass


# =================================================================================
# 14.4 — CLASS: InvoiceDetailsDialog
# =================================================================================
class InvoiceDetailsDialog:

    def __init__(self, page, db, invoice):
        self.page = page
        self.db = db
        self.invoice = invoice
        self.dialog = None
        self.setup_ui()

    def _fmt(self, d):
        if not d:
            return ''
        s = str(d).strip()
        if len(s) >= 10 and s[4] == '-' and s[7] == '-':
            return f"{s[8:10]}/{s[5:7]}/{s[0:4]}"
        return s[:10]

    def _f(self, v, d=Decimal('0')):
        try:
            return Decimal(str(v))
        except (InvalidOperation, ValueError, TypeError):
            return d

    def setup_ui(self):
        inv = self.invoice
        tid = inv.get('traveler_id', '')
        traveler = self.db.get_traveler_by_id(tid) if tid else None
        traveler_name = (
            f"{traveler.get('first_name', '')} "
            f"{traveler.get('last_name', '')}").strip() if traveler else 'Unknown'

        company_name = "Alhudha Haj Travel"
        try:
            if not self.db.company_settings.empty:
                company_name = str(
                    self.db.company_settings.iloc[0].get(
                        'company_name', company_name))
        except Exception:
            pass

        base = self._f(inv.get('amount', 0))
        disc_pct = self._f(inv.get('discount_percentage', 0))
        disc_amt = self._f(inv.get('discount_amount', 0))
        taxable = self._f(inv.get('taxable_value', base - disc_amt))
        gst_pct = self._f(inv.get('gst_percentage', 18))
        gst_amt = self._f(inv.get('gst_amount', 0))
        tcs_pct = self._f(inv.get('tcs_percentage', 0.1))
        tcs_amt = self._f(inv.get('tcs_amount', 0))
        total = self._f(inv.get('total_amount', 0))
        rounded = self._f(inv.get('rounded_total', total))

        try:
            words = number_to_words_indian(int(rounded))
        except Exception:
            words = ''

        def row(label, value, color=None):
            return ft.Row([
                ft.Text(label, width=180,
                        weight=ft.FontWeight.BOLD, size=11),
                ft.Text(str(value), size=11, expand=True, selectable=True,
                        color=color),
            ], spacing=8)

        content = ft.Column(
            controls=[
                ft.Text(company_name, size=15,
                        weight=ft.FontWeight.BOLD, color="#1e40af"),
                ft.Text(f"Invoice: {inv.get('invoice_no', '')}",
                        size=12, weight=ft.FontWeight.BOLD),
                ft.Divider(height=8),
                row("Date", self._fmt(inv.get('issue_date', ''))),
                row("Due Date", self._fmt(inv.get('due_date', ''))),
                row("Status", str(inv.get('status', 'pending')).upper()),
                ft.Divider(height=8),
                row("Traveler", traveler_name),
                ft.Divider(height=8),
                row("Package Cost", f"₹{base:,.2f}"),
                row(f"Discount @ {disc_pct:.2f}%",
                    f"- ₹{disc_amt:,.2f}", "#c2185b"),
                row("Taxable Value", f"₹{taxable:,.2f}"),
                row(f"GST @ {gst_pct}%", f"₹{gst_amt:,.2f}"),
                row("Total Before TCS",
                    f"₹{taxable + gst_amt:,.2f}"),
                row(f"TCS @ {tcs_pct}%", f"₹{tcs_amt:,.2f}"),
                ft.Divider(height=8),
                row("Total Amount", f"₹{total:,.2f}"),
                row("★ Rounded Total", f"₹{rounded:,.0f}", "#e67e22"),
                ft.Divider(height=8),
                ft.Text(f"Amount in Words: {words}", size=11,
                        italic=True, color="#27ae60"),
                ft.Text(f"Notes: {inv.get('notes', '')}", size=11,
                        color=ft.Colors.GREY_600),
            ], spacing=6, scroll=ft.ScrollMode.AUTO)

        self.dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text(f"📄 {inv.get('invoice_no', 'Invoice')}",
                          weight=ft.FontWeight.BOLD, size=14),
            content=ft.Container(content=content,
                                 width=520, height=560, padding=10),
            actions=[
                ft.Button(content=ft.Text("Close"),
                          on_click=lambda e: self.page.pop_dialog(),
                          bgcolor="#3498db", color=ft.Colors.WHITE),
            ],
            actions_alignment=ft.MainAxisAlignment.END,
        )

    def show(self):
        self.page.show_dialog(self.dialog)


# =================================================================================
# 14.5 — CLASS: ManualInvoiceDialog
# =================================================================================
class ManualInvoiceDialog:

    def __init__(self, page, db, current_user, on_save=None):
        self.page = page
        self.db = db
        self.current_user = current_user
        self.on_save = on_save
        self.settings_manager = SettingsManager(db)

        self.cust_name = None
        self.cust_phone = None
        self.cust_email = None
        self.traveler_dd = None
        self.package_dd = None
        self.items_table = None
        self.item_rows = []
        self.traveler_map = {}
        self._last_applied_tid = None
        self._selection_watch_running = False

        self.sub_lbl = None
        self.disc_lbl = None
        self.taxable_lbl = None
        self.total_lbl = None
        self.dup_warning_lbl = None
        self.gen_btn = None

        self.dialog = None
        self._typing_until = 0.0
        self.setup_ui()

    def _field(self, **kwargs):
        if "content_padding" not in kwargs:
            kwargs["content_padding"] = ft.Padding.symmetric(
                horizontal=10, vertical=10)
        field = ft.TextField(**kwargs)

        def _on_change(e, _f=field):
            self._typing_until = time.time() + TYPING_GUARD_SECONDS
            try:
                ctrl = getattr(e, "control", None)
                if ctrl is not None and hasattr(ctrl, "value"):
                    if _f.value != ctrl.value:
                        _f.value = ctrl.value
            except Exception:
                pass

        field.on_change = _on_change
        return field

    def setup_ui(self):
        topts = [ft.dropdown.Option(key="", text="Select Traveler")]
        self.traveler_map = {}
        try:
            for i, t in enumerate(self.db.get_travelers()):
                if not t.get('id'):
                    continue
                name = (f"{t.get('first_name', '')} "
                        f"{t.get('last_name', '')}").strip()
                key = f"t_{i}"
                self.traveler_map[key] = {
                    'id': str(t['id']).strip(),
                    'text': name
                }
                topts.append(ft.dropdown.Option(
                    key=key, text=name))
        except Exception:
            pass
        self.traveler_dd = ft.Dropdown(
            label="Traveler", options=topts, value="",
            height=48, text_size=12)
        
        self.traveler_dd.on_change = self.on_traveler_change

        self.dup_warning_lbl = ft.Text("", color=ft.Colors.RED_600, weight=ft.FontWeight.BOLD, size=11)

        bopts = [ft.dropdown.Option(key="", text="Select Batch")]
        try:
            for b in self.db.get_batches():
                bid = str(b.get('id') or '').strip()
                if not bid:
                    continue
                bopts.append(ft.dropdown.Option(
                    key=bid,
                    text=(f"{b.get('batch_name', '')} "
                          f"(₹{b.get('price', 0):,.0f})")))
        except Exception:
            pass
        self.package_dd = ft.Dropdown(
            label="Package", options=bopts, value="",
            height=48, text_size=12)

        self.cust_name = self._field(
            label="Customer Name *", height=48, text_size=12)
        self.cust_phone = self._field(
            label="Phone *", height=48, text_size=12)
        self.cust_email = self._field(
            label="Email", height=48, text_size=12)

        self.items_table = ft.DataTable(
            columns=[
                ft.DataColumn(ft.Text("#", size=11)),
                ft.DataColumn(ft.Text("Item", size=11)),
                ft.DataColumn(ft.Text("Rate", size=11)),
                ft.DataColumn(ft.Text("Qty", size=11)),
                ft.DataColumn(ft.Text("Amount", size=11)),
                ft.DataColumn(ft.Text("Disc%", size=11)),
                ft.DataColumn(ft.Text("Taxable", size=11)),
                ft.DataColumn(ft.Text("Total", size=11)),
                ft.DataColumn(ft.Text("", size=11)),
            ],
            rows=[], heading_row_color=ft.Colors.BLUE_GREY_800,
            heading_row_height=36, data_row_min_height=40,
            column_spacing=8,
        )

        self.sub_lbl = ft.Text("Package Cost: ₹0.00", size=12,
                               weight=ft.FontWeight.BOLD)
        self.disc_lbl = ft.Text("Discount: ₹0.00", size=12,
                                color="#c2185b",
                                weight=ft.FontWeight.BOLD)
        self.taxable_lbl = ft.Text("Taxable: ₹0.00", size=12,
                                   weight=ft.FontWeight.BOLD)
        self.total_lbl = ft.Text("★ Rounded Total: ₹0", size=14,
                                 weight=ft.FontWeight.BOLD,
                                 color="#e67e22")

        def add_item(e):
            bid = self.package_dd.value
            if not bid:
                self._snack("⚠️ Select a Package first")
                return
            try:
                batch = None
                for b in self.db.get_batches():
                    if str(b.get('id') or '').strip() == str(bid).strip():
                        batch = b
                        break
                if not batch:
                    return
                rate = _money(batch.get('price', 0) or 0)
                self.item_rows.append({
                    'item': batch.get('batch_name', 'Package'),
                    'hsn': '998555',
                    'rate': rate,
                    'qty': 1,
                    'discount_percent': 0,
                })
                self.refresh_items()
            except Exception as ex:
                self._snack(f"❌ {ex}")

        add_btn = ft.Button(
            content=ft.Text("➕ Add Item"),
            on_click=add_item, height=42,
            bgcolor="#3498db", color=ft.Colors.WHITE)

        def generate(e):
            self.generate_invoice()

        self.gen_btn = ft.Button(
            content=ft.Text("💾 Generate Invoice"),
            on_click=generate, height=42,
            bgcolor="#27ae60", color=ft.Colors.WHITE)

        content = ft.Column(
            controls=[
                ft.Text("🧾 Manual Invoice Creator", size=15,
                        weight=ft.FontWeight.BOLD, color="#1e40af"),
                ft.ResponsiveRow(
                    controls=[
                        ft.Container(content=self.traveler_dd,
                                     col={"xs": 12, "sm": 6, "md": 6}),
                    ], spacing=8, run_spacing=8),
                self.dup_warning_lbl,
                ft.ResponsiveRow(
                    controls=[
                        ft.Container(content=self.cust_name,
                                     col={"xs": 12, "sm": 4, "md": 4}),
                        ft.Container(content=self.cust_phone,
                                     col={"xs": 6, "sm": 4, "md": 4}),
                        ft.Container(content=self.cust_email,
                                     col={"xs": 6, "sm": 4, "md": 4}),
                    ], spacing=8, run_spacing=8),
                ft.Divider(),
                ft.ResponsiveRow(
                    controls=[
                        ft.Container(content=self.package_dd,
                                     col={"xs": 12, "sm": 8, "md": 9}),
                        ft.Container(content=add_btn,
                                     col={"xs": 12, "sm": 4, "md": 3}),
                    ], spacing=8, run_spacing=8),
                ft.Container(
                    content=ft.Row([self.items_table],
                                   scroll=ft.ScrollMode.ADAPTIVE),
                    bgcolor=ft.Colors.GREY_50,
                    border=ft.Border.all(1, ft.Colors.GREY_300),
                    border_radius=8, padding=5, height=240),
                ft.Divider(),
                self.sub_lbl, self.disc_lbl,
                self.taxable_lbl, self.total_lbl,
            ],
            spacing=10,
            scroll=ft.ScrollMode.AUTO,
        )

        self.dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("🧾 Manual Invoice",
                          weight=ft.FontWeight.BOLD, size=15),
            content=ft.Container(content=content,
                                 width=760, height=680, padding=10),
            actions=[
                ft.TextButton(content=ft.Text("Cancel"),
                              on_click=lambda e: self.page.pop_dialog()),
                self.gen_btn,
            ],
            actions_alignment=ft.MainAxisAlignment.END,
        )

    def on_traveler_change(self, e):
        control = getattr(e, 'control', None)
        event_val = (getattr(control, 'value', None)
                     or getattr(e, 'data', None)
                     or '')
        print(f"[ManualInvoice] traveler dropdown event: value={event_val!r}",
              flush=True)
        tid = self._resolve_tid(event_val)
        if tid:
            self._last_applied_tid = tid
            self._apply_traveler_change(tid)

        async def _process():
            for _ in range(15):
                await asyncio.sleep(0.1)
                for val in (getattr(self.traveler_dd, 'value', None),
                            getattr(control, 'value', None),
                            getattr(e, 'data', None), event_val):
                    candidate = self._resolve_tid(val)
                    if candidate and candidate != self._last_applied_tid:
                        self._last_applied_tid = candidate
                        self._apply_traveler_change(candidate)
                        return

        try:
            self.page.run_task(_process)
        except Exception as ex:
            print(f"[ManualInvoice] polling task failed: {ex}")

    def _resolve_tid(self, val):
        if not val:
            return None
        val_str = str(val).strip()
        
        # 1. Check map by key
        if val_str in self.traveler_map:
            return self.traveler_map[val_str]['id']
            
        # 2. Check map by exact text
        for k, info in self.traveler_map.items():
            if info['text'] == val_str:
                return info['id']
                
        # 3. Check database directly
        for t in self.db.get_travelers():
            name = f"{t.get('first_name', '')} {t.get('last_name', '')}".strip()
            if name == val_str or str(t['id']) == val_str:
                return str(t['id']).strip()
                
            # 4. Fuzzy match
            if val_str in name:
                 return str(t['id']).strip()
                 
        return None

    def _apply_traveler_change(self, tid):
        if not tid:
            self.cust_name.value = ""
            self.cust_phone.value = ""
            self.cust_email.value = ""
            self.package_dd.value = ""
            self.dup_warning_lbl.value = ""
            self.gen_btn.disabled = False
            self._force_ui_update()
            return

        traveler = self.db.get_traveler_by_id(tid)
        if traveler:
            self.cust_name.value = f"{traveler.get('first_name', '')} {traveler.get('last_name', '')}".strip()
            self.cust_phone.value = traveler.get('mobile', '')
            self.cust_email.value = traveler.get('email', '')
            
            bid = str(traveler.get('batch_id') or '').strip()
            if bid:
                self.package_dd.value = bid
            else:
                self.package_dd.value = ""

            existing_invoices = self.db.get_invoices(traveler_id=tid)
            if existing_invoices:
                inv_no = existing_invoices[0].get('invoice_no', '')
                self.dup_warning_lbl.value = f"⚠️ Notice: Invoice {inv_no} already exists for this traveler!"
                self.gen_btn.disabled = True
            else:
                self.dup_warning_lbl.value = ""
                self.gen_btn.disabled = False

        self._force_ui_update()

    async def _watch_traveler_selection(self):
        """Fallback for a browser dropdown that changes visually without
        consistently dispatching Flet on_change."""
        if self._selection_watch_running:
            return
        self._selection_watch_running = True
        print("[ManualInvoice] traveler watcher started", flush=True)
        try:
            for _ in range(240):
                await asyncio.sleep(0.25)
                tid = self._resolve_tid(getattr(self.traveler_dd, 'value', None))
                if tid and tid != self._last_applied_tid:
                    self._last_applied_tid = tid
                    self._apply_traveler_change(tid)
        except Exception as ex:
            print(f"[ManualInvoice] traveler watcher failed: {ex}")
        finally:
            self._selection_watch_running = False

    def _force_ui_update(self):
        try:
            self.dialog.update()
        except Exception:
            pass

    @staticmethod
    def _money(value):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except (InvalidOperation, ValueError, TypeError):
            return Decimal("0.00")

    @staticmethod
    def _round_rule(value):
        value = Decimal(str(value))
        floor_value = value.to_integral_value(rounding='ROUND_FLOOR')
        return (value.to_integral_value(rounding='ROUND_CEILING')
                if value - floor_value >= Decimal("0.50") else floor_value)

    def refresh_items(self):
        self.items_table.rows.clear()
        for i, item in enumerate(self.item_rows):
            rate = item['rate']
            qty = item['qty']
            amount = rate * Decimal(str(qty))
            disc_amt = amount * (Decimal(str(item['discount_percent'])) / Decimal('100'))
            taxable = amount - disc_amt
            
            tax = self.settings_manager.get_tax_settings()
            try:
                gst = _rate(tax.get('gst_percentage', 18))
            except Exception:
                gst = Decimal('18')
            try:
                tcs = _rate(tax.get('tcs_percentage', 0.1))
            except Exception:
                tcs = Decimal('0.1')

            gst_amt = taxable * gst / Decimal('100')
            taxable_with_gst = taxable + gst_amt
            tcs_amt = taxable_with_gst * tcs / Decimal('100')
            total = taxable_with_gst + tcs_amt

            def rm(idx=i):
                def handler(e):
                    self.item_rows.pop(idx)
                    self.refresh_items()
                return handler

            self.items_table.rows.append(
                ft.DataRow(cells=[
                    ft.DataCell(ft.Text(str(i + 1), size=10)),
                    ft.DataCell(ft.Text(item['item'][:20], size=10)),
                    ft.DataCell(ft.Text(f"₹{rate:,.0f}", size=10)),
                    ft.DataCell(ft.Text(str(qty), size=10)),
                    ft.DataCell(ft.Text(f"₹{amount:,.0f}", size=10)),
                    ft.DataCell(ft.Text(
                        f"{item['discount_percent']:.1f}%", size=10)),
                    ft.DataCell(ft.Text(f"₹{taxable:,.0f}", size=10)),
                    ft.DataCell(ft.Text(f"₹{total:,.0f}", size=10)),
                    ft.DataCell(ft.IconButton(
                        icon=ft.Icons.DELETE, icon_size=14,
                        icon_color="#e74c3c",
                        on_click=rm())),
                ]))

        tax = self.settings_manager.get_tax_settings()
        try:
            gst = _rate(tax.get('gst_percentage', 18))
            tcs = _rate(tax.get('tcs_percentage', 0.1))
        except Exception:
            gst = Decimal('18')
            tcs = Decimal('0.1')

        sub = sum((i['rate'] * Decimal(str(i['qty'])) for i in self.item_rows), Decimal('0.00'))
        disc = sum((i['rate'] * Decimal(str(i['qty'])) * (Decimal(str(i['discount_percent'])) / Decimal('100'))
                   for i in self.item_rows), Decimal('0.00'))
        taxable = sub - disc
        gst_amt = taxable * gst / Decimal('100')
        taxable_with_gst = taxable + gst_amt
        tcs_amt = taxable_with_gst * tcs / Decimal('100')
        total = taxable_with_gst + tcs_amt
        rounded = _round_rule(total)

        self.sub_lbl.value = f"Package Cost: ₹{sub:,.2f}"
        self.disc_lbl.value = f"Discount: ₹{disc:,.2f}"
        self.taxable_lbl.value = f"Taxable: ₹{taxable:,.2f}"
        self.total_lbl.value = f"★ Rounded Total: ₹{rounded:,.0f}"
        
        for c in (self.items_table, self.sub_lbl, self.disc_lbl,
                  self.taxable_lbl, self.total_lbl):
            try:
                c.update()
            except Exception:
                pass

    def generate_invoice(self):
        if not self.item_rows:
            self._snack("⚠️ Add at least one item")
            return
        name = (self.cust_name.value or '').strip()
        phone = (self.cust_phone.value or '').strip()
        if not name or not phone:
            self._snack("⚠️ Customer name and phone required")
            return

        val = getattr(self.traveler_dd, 'value', None)
        tid = (getattr(self, '_last_applied_tid', None)
               or self._resolve_tid(val))

        bid = str(getattr(self.package_dd, 'value', None) or '').strip()

        if tid:
            traveler = self.db.get_traveler_by_id(tid)
            if traveler and bid:
                t_batch = str(traveler.get('batch_id') or '').strip()
                if t_batch and t_batch != bid:
                    self._snack("⚠️ Selected package does not match the traveler's batch.")
                    return

            existing_invoices = self.db.get_invoices(traveler_id=tid)
            if existing_invoices:
                self._snack(f"⚠️ Invoice {existing_invoices[0].get('invoice_no')} already exists for this traveler.")
                return

        try:
            if not tid:
                parts = name.split()
                traveler_data = {
                    'first_name': parts[0] if parts else '',
                    'last_name': ' '.join(parts[1:]) if len(parts) > 1 else '',
                    'passport_name': name.upper(),
                    'mobile': phone,
                    'email': (self.cust_email.value or '').strip(),
                    'passport_no': 'N/A',
                    'status': 'Active',
                    'batch_id': bid}
                tid = self.db.add_traveler(traveler_data)

            tax = self.settings_manager.get_tax_settings()
            gst = _rate(tax.get('gst_percentage', 18))
            tcs = _rate(tax.get('tcs_percentage', 0.1))

            sub = sum((i['rate'] * Decimal(str(i['qty'])) for i in self.item_rows), Decimal('0.00'))
            disc = sum((i['rate'] * Decimal(str(i['qty'])) * (Decimal(str(i['discount_percent'])) / Decimal('100'))
                       for i in self.item_rows), Decimal('0.00'))

            inv_id = self.db.generate_invoice(
                tid, bid, sub, gst, tcs,
                discount_percentage=Decimal('0'),
                discount_amount=disc)

            try:
                self.db.log_activity(
                    self.current_user['id'], "manual_invoice",
                    f"Created manual invoice for {name}")
            except Exception:
                pass

            self._typing_until = 0.0
            self.page.pop_dialog()
            self._snack(f"✅ Invoice generated: {inv_id}")
            if self.on_save:
                self.on_save()
        except Exception as ex:
            import traceback
            traceback.print_exc()
            self._snack(f"❌ {ex}")

    def show(self):
        self.page.show_dialog(self.dialog)
        try:
            self.page.run_task(self._watch_traveler_selection)
        except Exception as ex:
            print(f"[ManualInvoice] traveler watcher could not start: {ex}")

    def _snack(self, msg):
        try:
            self.page.show_dialog(ft.SnackBar(content=ft.Text(msg)))
        except Exception:
            pass


# =================================================================================
# SECTION 14 END (FLET 1.0.0 VERSION) — v4.5
# =================================================================================
